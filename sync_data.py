"""@description Excel·JSON 양방향 동기화"""

import json
import argparse
import sys
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Optional, Tuple
from copy import deepcopy

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.styles.colors import Color

from csat_benchmark.answers import CorrectAnswer, _is_correct_answer, _validate_correct_answer
from csat_benchmark.evaluation import EvaluationError, resolve_run
from csat_benchmark.configuration import load_config
from csat_benchmark.exams import load_exam
from csat_benchmark.exports import (
    export_run_to_excel,
    import_excel_corrections,
    publish_run,
)
from csat_benchmark.metadata import sync_model_metadata


REFUSAL_ANSWER = -2
NO_ANSWER = -1
REFUSAL_MARKERS = {"-2", "(검열)", "검열", "Refusal", "refusal"}
DEFAULT_EXCEL_PATH = Path('2026 수능 LLM 풀이.xlsx')
DEFAULT_HARD_EXCEL_PATH = Path('2026 수능 LLM 풀이 hard.xlsx')
DEFAULT_CONFIG_PATH = Path(__file__).resolve().parent / 'config.json'
DEFAULT_MODEL_METADATA_PATH = Path(__file__).resolve().parent / 'web' / 'model_metadata.json'


def _config_project_root(config_path: Path) -> Path:
    """@description 설정 파일 기준 저장소·작업 루트 탐색"""
    for parent in (config_path.parent, *config_path.parent.parents):
        if (parent / '.env').is_file() or (parent / '.git').exists():
            return parent
    return config_path.parent


def _metadata_path_for_config(config_path: str | Path) -> Path:
    """@description 설정 파일 기준 모델 메타데이터 경로 반환"""
    resolved = Path(config_path).expanduser().resolve()
    return _config_project_root(resolved) / 'web' / 'model_metadata.json'


def _generic_sync_main(arguments: List[str]) -> int:
    """@description 시험·쉬움 모드 기준 공개·Excel 동기화 처리"""
    parser = argparse.ArgumentParser(
        description="시험 실행 결과 공개·Excel 동기화",
        allow_abbrev=False,
    )
    subparsers = parser.add_subparsers(dest="command", required=True)
    for command, help_text in (
        ("publish", "완료된 실행을 공개 경로에 병합"),
        ("import", "비공개 검증 결과를 Excel로 가져오기"),
        ("export", "Excel 수동 답안을 비공개 검증 결과에 반영"),
    ):
        command_parser = subparsers.add_parser(command, help=help_text, allow_abbrev=False)
        command_parser.add_argument("--exam", required=True, help="시험 ID 또는 매니페스트 경로")
        command_parser.add_argument("--easy", action="store_true", help="쉬움 모드 선택")
        command_parser.add_argument("--excel", help="Excel 파일 경로")
        command_parser.add_argument("--output", help="verified 결과 출력 경로")
        command_parser.add_argument("--published-dir", help="공개 산출물 경로 재정의")
        command_parser.add_argument("--config", help="공개 메타데이터 설정 경로")
        command_parser.add_argument("--models", nargs="+", help="대상 모델 이름")
        command_parser.add_argument("--targets", nargs="+", help="시험 target 목록")
        command_parser.add_argument("--subject", help="한 섹션을 고를 과목")
        command_parser.add_argument("--section", help="한 섹션을 고를 영역")
        command_parser.add_argument("--subjects", nargs="+", help="과목 또는 탐구 영역 목록")
        command_parser.add_argument("--question-numbers", nargs="+", type=int, help="문항 번호")
        command_parser.add_argument("--benchmark-all", action="store_true", help="시험의 모든 섹션 선택")
    args = parser.parse_args(arguments)
    if (args.subject is None) != (args.section is None):
        parser.error("--subject와 --section은 함께 지정해야 합니다.")
    if args.command == "publish" and args.output:
        parser.error("publish 명령에는 --output을 사용할 수 없습니다.")
    if args.command == "publish" and args.excel:
        parser.error("publish 명령에는 --excel을 사용할 수 없습니다.")
    if args.command != "publish" and args.published_dir:
        parser.error(f"{args.command} 명령에는 --published-dir를 사용할 수 없습니다.")
    try:
        exam_argument = Path(args.exam).expanduser()
        exam_input = (
            load_exam(exam_argument, project_root=exam_argument.resolve().parent)
            if exam_argument.is_file()
            else args.exam
        )
        exam, _mode, run, run_path = resolve_run(
            exam_input,
            easy=args.easy,
            model_names=args.models,
        )
        excel_path = (
            Path(args.excel).expanduser().resolve()
            if args.excel
            else exam.project_root / (DEFAULT_EXCEL_PATH if args.easy else DEFAULT_HARD_EXCEL_PATH)
        )
        common = {
            "model_names": args.models,
            "targets": args.targets,
            "subject": args.subject,
            "section": args.section,
            "subjects": args.subjects,
            "benchmark_all": args.benchmark_all,
            "question_numbers": args.question_numbers,
        }
        if args.command == "publish":
            paths = publish_run(
                exam,
                run_path,
                run_path,
                output_dir=args.published_dir,
                config_path=args.config,
                **common,
            )
            for name, path in paths.items():
                print(f"{name}: {path}")
            return 0
        if args.command == "import":
            output = export_run_to_excel(
                run_path,
                excel_path,
                run=run_path,
                exam=exam,
                **common,
            )
            print(f"Excel 가져오기 완료: {output}")
            return 0
        result = import_excel_corrections(
            run_path,
            excel_path,
            output_path=args.output or run_path,
            run=run_path,
            exam=exam,
            **common,
        )
        print(f"Excel 수동 답안 반영 완료: {result['path']} (변경 {result['changed']}건)")
        return 0
    except (OSError, ValueError, EvaluationError) as error:
        print(f"오류: {error}", file=sys.stderr)
        return 2


def _sync_model_metadata(model_config: Dict[str, Dict],
                         metadata_path: Path = DEFAULT_MODEL_METADATA_PATH) -> bool:
    """
    @description 모델 설정·시각 입력 지원 여부 대시보드 메타데이터 동기화

    @param model_config 모델 이름별 설정 매핑
    @param metadata_path 대시보드 메타데이터 파일 경로
    @return 파일 내용 변경 여부
    """
    return sync_model_metadata(model_config, metadata_path)


def normalize_answer_value(answer):
    """@description 답안 값 → JSON 채점용 숫자·상태 변환"""
    if answer is None or answer == '' or str(answer).strip() == '':
        return NO_ANSWER, 'no_answer'

    answer_text = str(answer).strip()
    if answer == NO_ANSWER or answer_text in {'-1', '(포기)'}:
        return NO_ANSWER, 'no_answer'
    if answer_text in REFUSAL_MARKERS:
        return REFUSAL_ANSWER, 'refusal'

    try:
        return int(answer), 'answered'
    except (ValueError, TypeError):
        return NO_ANSWER, 'parse_failed'


def _normalise_correct_answer_cell(value) -> Optional[CorrectAnswer]:
    """@description Excel 정답 셀을 단일값·복수값으로 변환"""
    if value is None or str(value).strip() == '' or isinstance(value, bool):
        return None
    if isinstance(value, str):
        parts = [part.strip() for part in value.strip().split(',')]
        if any(not part for part in parts):
            return None
        try:
            parsed = [int(part) for part in parts] if len(parts) > 1 else int(parts[0])
        except (ValueError, TypeError):
            return None
    else:
        parsed = value
    try:
        return _validate_correct_answer(parsed)
    except (TypeError, ValueError):
        return None


def _format_correct_answer_cell(value: CorrectAnswer) -> int | str:
    """@description 정답 단일값·복수값을 Excel 셀 값으로 변환"""
    if isinstance(value, list):
        return ', '.join(str(answer) for answer in value)
    return value


def _create_hard_excel_template(source_path: Path, target_path: Path):
    """@description 정답 구조 기반 hard용 Excel 템플릿 생성"""
    if not source_path.exists():
        raise FileNotFoundError(f"hard Excel 템플릿 원본을 찾을 수 없습니다: {source_path}")

    workbook = load_workbook(source_path)
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        header_row = None
        for row_idx in range(1, min(6, ws.max_row + 1)):
            cell_value = ws.cell(row=row_idx, column=1).value
            if cell_value and '문항 번호' in str(cell_value):
                header_row = row_idx
                break

        if not header_row:
            continue

        answer_col = 2
        for col_idx in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=header_row, column=col_idx).value
            if cell_value and str(cell_value).strip() == '정답':
                answer_col = col_idx
                break

        if ws.max_column > answer_col:
            ws.delete_cols(answer_col + 1, ws.max_column - answer_col)

    target_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target_path)
    print(f"hard Excel 템플릿 생성 완료: {target_path}")


class PathMapper:
    """@description Excel 시트명·JSON 경로 매핑"""

    # 시트명 → JSON 경로 매핑
    SHEET_TO_JSON = {
        '국어-공통': 'problems/국어/공통',
        '국어-화작': 'problems/국어/화작',
        '국어-언매': 'problems/국어/언매',
        '수학-공통': 'problems/수학/공통',
        '수학-확통': 'problems/수학/확통',
        '수학-미적': 'problems/수학/미적',
        '수학-기하': 'problems/수학/기하',
        '영어': 'problems/영어',
        '한국사': 'problems/한국사',
        '물리1': 'problems/탐구/물리1',
        '화학1': 'problems/탐구/화학1',
        '생명1': 'problems/탐구/생명1',
        '사회문화': 'problems/탐구/사회문화',
    }

    def __init__(self, base_dir: Path = None):
        self.base_dir = base_dir or Path('.')
        # 역방향 매핑 생성
        self._json_to_sheet = {v: k for k, v in self.SHEET_TO_JSON.items()}

    def sheet_to_json_path(self, sheet_name: str) -> Optional[Path]:
        """@description 시트 이름 기준 JSON 폴더 경로 반환"""
        if sheet_name in self.SHEET_TO_JSON:
            return self.base_dir / self.SHEET_TO_JSON[sheet_name]
        return None

    def json_to_sheet_name(self, json_path: Path) -> Optional[str]:
        """@description JSON 경로 기준 시트 이름 반환"""
        # 경로 정규화
        rel_path = str(json_path).replace('\\', '/')
        # results_verified.json 제거
        if rel_path.endswith('/results_verified.json'):
            rel_path = rel_path[:-len('/results_verified.json')]
        elif rel_path.endswith('results_verified.json'):
            rel_path = str(Path(rel_path).parent).replace('\\', '/')

        # 매핑 검색
        for json_rel, sheet in self._json_to_sheet.items():
            if rel_path.endswith(json_rel) or json_rel in rel_path:
                return sheet
        return None

    def get_all_sheets(self) -> List[str]:
        """@description 전체 시트 이름 반환"""
        return list(self.SHEET_TO_JSON.keys())

    def get_subject_section(self, sheet_name: str) -> Tuple[str, str]:
        """@description 시트 이름 기준 과목·섹션 추출"""
        if '-' in sheet_name:
            parts = sheet_name.split('-', 1)
            return parts[0].strip(), parts[1].strip()
        return sheet_name.strip(), sheet_name.strip()


class ModelNameMapper:
    """@description JSON·Excel 모델 이름 매핑"""

    def __init__(self, mapping_file: Path = None):
        self.mapping_file = mapping_file or Path('model_mapping.json')
        self.mapping = {}  # JSON name -> Excel name
        self._load_mapping()

    def _load_mapping(self):
        """@description 매핑 파일 로드"""
        if self.mapping_file.exists():
            with open(self.mapping_file, 'r', encoding='utf-8') as f:
                self.mapping = json.load(f)

    def save_mapping(self):
        """@description 매핑 파일 저장"""
        with open(self.mapping_file, 'w', encoding='utf-8') as f:
            json.dump(self.mapping, f, ensure_ascii=False, indent=2)

    def json_to_excel(self, json_name: str) -> str:
        """@description JSON 모델 이름 → Excel 열 이름 변환"""
        return self.mapping.get(json_name, json_name)

    def excel_to_json(self, excel_name: str) -> str:
        """@description Excel 열 이름 → JSON 모델 이름 변환"""
        # 역방향 검색
        for json_name, mapped_excel in self.mapping.items():
            if mapped_excel == excel_name:
                return json_name
        return excel_name

    def add_mapping(self, json_name: str, excel_name: str):
        """@description 모델 이름 매핑 추가"""
        self.mapping[json_name] = excel_name


class ExcelHandler:
    """@description Excel 파일 읽기·쓰기"""

    def __init__(self, excel_path: Path):
        self.excel_path = Path(excel_path)
        self.workbook = None
        self._header_row_cache = {}

    def _load_workbook(self):
        """@description 첫 접근 시 통합 문서 로드"""
        if self.workbook is None:
            self.workbook = load_workbook(self.excel_path)

    def get_sheet_names(self) -> List[str]:
        """@description 전체 시트 이름 반환"""
        self._load_workbook()
        return self.workbook.sheetnames

    def _find_header_row(self, sheet_name: str) -> int:
        """@description 헤더 행 번호 조회 (1-based)"""
        if sheet_name in self._header_row_cache:
            return self._header_row_cache[sheet_name]

        self._load_workbook()
        ws = self.workbook[sheet_name]

        for row_idx in range(1, min(6, ws.max_row + 1)):
            cell_value = ws.cell(row=row_idx, column=1).value
            if cell_value and '문항 번호' in str(cell_value):
                self._header_row_cache[sheet_name] = row_idx
                return row_idx

        raise ValueError(f"'{sheet_name}' 시트에서 헤더 행을 찾을 수 없습니다.")

    def _find_score_row(self, sheet_name: str) -> int:
        """@description 총점 행 번호 조회 (1-based)"""
        self._load_workbook()
        ws = self.workbook[sheet_name]
        header_row = self._find_header_row(sheet_name)

        for row_idx in range(header_row + 1, ws.max_row + 1):
            cell_value = ws.cell(row=row_idx, column=1).value
            if cell_value and str(cell_value).strip() in ['총점', '총합', '점수']:
                return row_idx

        raise ValueError(f"'{sheet_name}' 시트에서 총점 행을 찾을 수 없습니다.")

    def get_model_columns(self, sheet_name: str) -> Dict[str, int]:
        """@description 모델 열 이름·열 번호 반환 (1-based)"""
        self._load_workbook()
        ws = self.workbook[sheet_name]
        header_row = self._find_header_row(sheet_name)

        models = {}
        for col_idx in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=header_row, column=col_idx).value
            if cell_value:
                col_str = str(cell_value).strip()
                # 불필요한 열 제외
                if col_str in ['문항 번호', '정답', 'nan', '']:
                    continue
                if 'Unnamed' in col_str:
                    continue
                models[col_str] = col_idx

        return models

    def get_model_answers(self, sheet_name: str, model_name: str) -> Dict[int, any]:
        """@description 모델별 문항 답 추출"""
        self._load_workbook()
        ws = self.workbook[sheet_name]
        header_row = self._find_header_row(sheet_name)
        score_row = self._find_score_row(sheet_name)

        model_columns = self.get_model_columns(sheet_name)
        if model_name not in model_columns:
            raise ValueError(f"'{model_name}' 모델을 '{sheet_name}' 시트에서 찾을 수 없습니다.")

        col_idx = model_columns[model_name]
        answers = {}

        for row_idx in range(header_row + 1, score_row):
            q_num = ws.cell(row=row_idx, column=1).value
            answer = ws.cell(row=row_idx, column=col_idx).value

            if q_num is not None:
                try:
                    q_num = int(q_num)
                    answers[q_num] = answer
                except (ValueError, TypeError):
                    pass

        return answers

    def calculate_score_from_answers(self, sheet_name: str, answers: Dict[int, any]) -> int:
        """@description 문항 답 기준 총점 계산"""
        self._load_workbook()
        correct_answers = self._get_correct_answers(sheet_name)
        questions_data = self._load_questions_for_sheet(sheet_name)
        points_by_question = {q['number']: q['points'] for q in questions_data['questions']}

        score = 0
        for q_num, correct in correct_answers.items():
            answer = answers.get(q_num)
            normalized_answer, _ = normalize_answer_value(answer)
            if _is_correct_answer(normalized_answer, correct):
                score += points_by_question.get(q_num, 0)

        return score

    def _load_questions_for_sheet(self, sheet_name: str) -> Dict:
        """@description 시트 대응 questions.json 로드"""
        json_path = PathMapper(base_dir=Path('.')).sheet_to_json_path(sheet_name)
        if not json_path:
            raise ValueError(f"'{sheet_name}' 시트에 대한 경로 매핑을 찾을 수 없습니다.")

        questions_file = json_path / 'questions.json'
        if not questions_file.exists():
            raise FileNotFoundError(f"questions.json을 찾을 수 없습니다: {questions_file}")

        with open(questions_file, 'r', encoding='utf-8') as f:
            return json.load(f)

    def get_model_score(self, sheet_name: str, model_name: str) -> Optional[int]:
        """@description 특정 모델 총점 반환"""
        self._load_workbook()
        ws = self.workbook[sheet_name]
        score_row = self._find_score_row(sheet_name)

        model_columns = self.get_model_columns(sheet_name)
        if model_name not in model_columns:
            return None

        col_idx = model_columns[model_name]
        score = ws.cell(row=score_row, column=col_idx).value

        try:
            return int(score)
        except (ValueError, TypeError):
            return None

    def get_max_score(self, sheet_name: str) -> int:
        """@description 정답 열 총점 기준 만점 반환"""
        self._load_workbook()
        ws = self.workbook[sheet_name]
        header_row = self._find_header_row(sheet_name)
        score_row = self._find_score_row(sheet_name)

        # 정답 열 찾기
        for col_idx in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=header_row, column=col_idx).value
            if cell_value and str(cell_value).strip() == '정답':
                max_score = ws.cell(row=score_row, column=col_idx).value
                try:
                    return int(max_score)
                except:
                    pass

        return 100  # 기본값

    def _find_answer_column(self, sheet_name: str) -> int:
        """@description 정답 열 번호 조회(1-based)"""
        self._load_workbook()
        ws = self.workbook[sheet_name]
        header_row = self._find_header_row(sheet_name)

        for col_idx in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=header_row, column=col_idx).value
            if cell_value and str(cell_value).strip() == '정답':
                return col_idx

        return 2  # 기본값

    def _get_correct_answers(self, sheet_name: str) -> Dict[int, any]:
        """@description 정답 데이터 추출({문항번호: 정답})"""
        self._load_workbook()
        ws = self.workbook[sheet_name]
        header_row = self._find_header_row(sheet_name)
        score_row = self._find_score_row(sheet_name)
        answer_col = self._find_answer_column(sheet_name)

        correct_answers = {}
        for row_idx in range(header_row + 1, score_row):
            q_num = ws.cell(row=row_idx, column=1).value
            correct = ws.cell(row=row_idx, column=answer_col).value
            if q_num is not None:
                try:
                    q_num = int(q_num)
                    normalized_correct = _normalise_correct_answer_cell(correct)
                    if normalized_correct is not None:
                        correct_answers[q_num] = normalized_correct
                except (ValueError, TypeError):
                    pass

        return correct_answers

    def update_correct_answer_cells(self, sheet_name: str, questions_data: Dict) -> None:
        """@description 문항 원본 기준 Excel 정답 셀 갱신"""
        self._load_workbook()
        ws = self.workbook[sheet_name]
        header_row = self._find_header_row(sheet_name)
        score_row = self._find_score_row(sheet_name)
        answer_col = self._find_answer_column(sheet_name)
        correct_answers = {
            int(question['number']): _validate_correct_answer(question['correct_answer'])
            for question in questions_data.get('questions', [])
        }

        for row_idx in range(header_row + 1, score_row):
            q_num = ws.cell(row=row_idx, column=1).value
            if isinstance(q_num, bool) or not isinstance(q_num, int):
                continue
            if q_num in correct_answers:
                ws.cell(row=row_idx, column=answer_col).value = _format_correct_answer_cell(
                    correct_answers[q_num]
                )

    def add_model_column(self, sheet_name: str, model_name: str,
                         answers: Dict[int, any], score: int,
                         position: Optional[int] = None,
                         after_model: Optional[str] = None) -> int:
        """
        @description 새 모델 열 추가
        @param sheet_name 시트 이름
        @param model_name 모델 이름
        @param answers {문항번호: 답} dict
        @param score 총점
        @param position 삽입 열 번호(1-based, None 시 마지막)
        @param after_model 삽입 기준 모델 이름
        @return 삽입 열 번호
        """
        self._load_workbook()
        ws = self.workbook[sheet_name]
        header_row = self._find_header_row(sheet_name)
        score_row = self._find_score_row(sheet_name)

        model_columns = self.get_model_columns(sheet_name)

        # 대상 모델 존재 여부 확인
        if model_name in model_columns:
            raise ValueError(f"'{model_name}' 모델이 이미 존재합니다. --update 옵션을 사용하세요.")

        # 삽입 위치 결정
        if after_model and after_model in model_columns:
            insert_col = model_columns[after_model] + 1
        elif position is not None:
            insert_col = position
        else:
            # 마지막 모델 열 다음
            if model_columns:
                insert_col = max(model_columns.values()) + 1
            else:
                insert_col = 3  # 정답 다음

        # 열 삽입
        ws.insert_cols(insert_col)

        # 스타일 정의
        bold_font = Font(bold=True)
        center_align = Alignment(horizontal='center', vertical='center')
        red_font = Font(color='FF0000')
        red_bold_font = Font(color='FF0000', bold=True)
        purple_font = Font(color='7C3AED')

        # 정답 데이터 조회
        correct_answers = self._get_correct_answers(sheet_name)

        # 헤더 설정 (볼드 + 중앙정렬)
        header_cell = ws.cell(row=header_row, column=insert_col, value=model_name)
        header_cell.font = bold_font
        header_cell.alignment = center_align

        # 답안 입력
        for row_idx in range(header_row + 1, score_row):
            q_num = ws.cell(row=row_idx, column=1).value
            if q_num is not None:
                try:
                    q_num = int(q_num)
                    cell = ws.cell(row=row_idx, column=insert_col)
                    cell.alignment = center_align

                    if q_num in answers:
                        answer = answers[q_num]
                        correct = correct_answers.get(q_num)

                        if answer == REFUSAL_ANSWER or str(answer).strip() in REFUSAL_MARKERS:
                            cell.value = "(검열)"
                            cell.font = purple_font
                        elif answer is None or answer == NO_ANSWER or answer == "":
                            cell.value = "(포기)"
                            cell.font = red_font
                        else:
                            cell.value = answer
                            normalized_answer, _ = normalize_answer_value(answer)
                            if not _is_correct_answer(normalized_answer, correct):
                                cell.font = red_font
                    else:
                        # answers 미포함 문항 포기 처리
                        cell.value = "(포기)"
                        cell.font = red_font

                except (ValueError, TypeError):
                    pass

        # 총점 입력 (중앙정렬)
        score_cell = ws.cell(row=score_row, column=insert_col, value=score)
        score_cell.alignment = center_align

        # 캐시 무효화
        self._header_row_cache.pop(sheet_name, None)

        return insert_col

    def update_model_column(self, sheet_name: str, model_name: str,
                            answers: Dict[int, any], score: int):
        """@description 모델 열 업데이트"""
        self._load_workbook()
        ws = self.workbook[sheet_name]
        header_row = self._find_header_row(sheet_name)
        score_row = self._find_score_row(sheet_name)

        model_columns = self.get_model_columns(sheet_name)
        if model_name not in model_columns:
            raise ValueError(f"'{model_name}' 모델을 찾을 수 없습니다.")

        col_idx = model_columns[model_name]

        # 스타일 정의
        bold_font = Font(bold=True)
        center_align = Alignment(horizontal='center', vertical='center')
        red_font = Font(color='FF0000')
        purple_font = Font(color='7C3AED')

        # 정답 데이터 조회
        correct_answers = self._get_correct_answers(sheet_name)

        # 헤더 스타일 (볼드 + 중앙정렬)
        header_cell = ws.cell(row=header_row, column=col_idx)
        header_cell.font = bold_font
        header_cell.alignment = center_align

        # 답안 업데이트
        for row_idx in range(header_row + 1, score_row):
            q_num = ws.cell(row=row_idx, column=1).value
            if q_num is not None:
                try:
                    q_num = int(q_num)
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.alignment = center_align
                    cell.font = Font()  # 기본 폰트 초기화

                    if q_num in answers:
                        answer = answers[q_num]
                        correct = correct_answers.get(q_num)

                        if answer == REFUSAL_ANSWER or str(answer).strip() in REFUSAL_MARKERS:
                            cell.value = "(검열)"
                            cell.font = purple_font
                        elif answer is None or answer == NO_ANSWER or answer == "":
                            cell.value = "(포기)"
                            cell.font = red_font
                        else:
                            cell.value = answer
                            normalized_answer, _ = normalize_answer_value(answer)
                            if not _is_correct_answer(normalized_answer, correct):
                                cell.font = red_font
                    else:
                        cell.value = "(포기)"
                        cell.font = red_font

                except (ValueError, TypeError):
                    pass

        # 총점 업데이트 (중앙정렬)
        score_cell = ws.cell(row=score_row, column=col_idx, value=score)
        score_cell.alignment = center_align

    def save(self):
        """@description 변경 사항 저장"""
        if self.workbook:
            self.workbook.save(self.excel_path)
            print(f"저장 완료: {self.excel_path}")


class DataConverter:
    """@description 데이터 변환"""

    def __init__(self, path_mapper: PathMapper, model_mapper: ModelNameMapper):
        self.path_mapper = path_mapper
        self.model_mapper = model_mapper

    def load_questions(self, sheet_name: str) -> Dict:
        """@description questions.json 로드"""
        json_path = self.path_mapper.sheet_to_json_path(sheet_name)
        if not json_path:
            raise ValueError(f"'{sheet_name}' 시트에 대한 경로 매핑을 찾을 수 없습니다.")

        questions_file = json_path / 'questions.json'
        if not questions_file.exists():
            raise FileNotFoundError(f"questions.json을 찾을 수 없습니다: {questions_file}")

        with open(questions_file, 'r', encoding='utf-8') as f:
            return json.load(f)

    def excel_to_json(self, sheet_name: str, model_name: str,
                      answers: Dict[int, any], excel_handler: ExcelHandler) -> Dict:
        """@description Excel 모델 열 데이터 → results_verified.json 형식 변환

        @param sheet_name 시트 이름
        @param model_name Excel 모델 이름
        @param answers {문항번호: 답}
        @param excel_handler ExcelHandler 인스턴스

        @return results_verified.json 형식 dict
        """
        questions_data = self.load_questions(sheet_name)
        subject, section = self.path_mapper.get_subject_section(sheet_name)
        json_model_name = self.model_mapper.excel_to_json(model_name)

        results = []
        total_score = 0
        correct_count = 0

        for q in questions_data['questions']:
            q_num = q['number']
            correct_answer = _validate_correct_answer(q['correct_answer'])
            points = q['points']

            extracted = answers.get(q_num)
            extracted_normalized, answer_status = normalize_answer_value(extracted)

            is_correct = _is_correct_answer(extracted_normalized, correct_answer)
            if is_correct:
                total_score += points
                correct_count += 1

            results.append({
                'question_number': q_num,
                'model_name': json_model_name,
                'extracted_answer': extracted_normalized,
                'correct_answer': correct_answer,
                'is_correct': is_correct,
                'points': points,
                'needs_manual_review': False,
                'answer_status': answer_status,
                'provider_stop_reason': 'refusal' if answer_status == 'refusal' else None,
                'raw_response': ''
            })

        total_points = sum(q['points'] for q in questions_data['questions'])

        return {
            'subject': questions_data.get('subject', subject),
            'section': questions_data.get('section', section),
            'timestamp': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'total_points': total_points,
            'total_verified': len(results),
            'correct_count': correct_count,
            'manual_review_count': 0,
            'model_scores': {json_model_name: total_score},
            'results': results
        }

    def json_to_excel(self, json_data: Dict, target_model: str = None) -> List[Tuple[str, Dict[int, any], int]]:
        """@description results_verified.json 데이터 → Excel 형식 변환

        @param json_data results_verified.json 데이터
        @param target_model 특정 모델 변환(None 시 전체 모델)

        @return [(model_name, {문항번호: 답}, 총점), ...] 목록
        """
        model_scores = json_data.get('model_scores', {})
        if not model_scores:
            raise ValueError("model_scores가 비어있습니다.")

        results_list = []

        for json_model_name, score in model_scores.items():
            # 특정 모델 처리
            if target_model and json_model_name != target_model:
                continue

            excel_model_name = self.model_mapper.json_to_excel(json_model_name)

            # 답안 추출
            answers = {}
            for result in json_data.get('results', []):
                if result['model_name'] == json_model_name:
                    q_num = result['question_number']
                    extracted = result['extracted_answer']
                    answer_status = result.get('answer_status')
                    if answer_status == 'refusal' or extracted == REFUSAL_ANSWER:
                        answers[q_num] = "(검열)"
                    elif extracted == NO_ANSWER:
                        answers[q_num] = None
                    else:
                        answers[q_num] = extracted

            results_list.append((excel_model_name, answers, score))

        return results_list


class SyncManager:
    """동기화 관리"""

    WRONG_ONLY_BASE_MODELS = {
        'GPT-5.6 Sol (max*)': 'GPT-5.6 Sol (high)',
        'GPT-5.6 Terra (max*)': 'GPT-5.6 Terra (high)',
        'GPT-5.6 Luna (max*)': 'GPT-5.6 Luna (high)',
        'GPT-5.5 (xhigh*)': 'GPT-5.5 (high)',
        'GPT-5.4 (xhigh*)': 'GPT-5.4 (high)',
        'Claude Opus 4.7 (max*)': 'Claude Opus 4.7 (high)',
    }

    def __init__(self, excel_path: Path, problems_dir: Path = None,
                 model_mapping_path: Path = None, hard_mode: bool = False):
        self.excel_path = Path(excel_path)
        self.problems_dir = problems_dir or Path('problems')
        self.hard_mode = hard_mode
        self.verified_filename = 'hard_results_verified.json' if hard_mode else 'results_verified.json'
        self.raw_results_filename = 'hard_results.json' if hard_mode else 'results.json'
        self.all_results_filename = 'hard_all_results.json' if hard_mode else 'all_results.json'
        self.token_usage_filename = 'hard_token_usage.json' if hard_mode else 'token_usage.json'

        self.path_mapper = PathMapper(base_dir=Path('.'))
        self.model_mapper = ModelNameMapper(model_mapping_path)
        self.excel_handler = ExcelHandler(self.excel_path)
        self.converter = DataConverter(self.path_mapper, self.model_mapper)
        self._token_usage = self._load_token_usage()
        self._model_config = self._load_model_config()
        self.model_metadata_path = _metadata_path_for_config(DEFAULT_CONFIG_PATH)

    def _load_model_config(self) -> Dict[str, Dict]:
        """@description 모델 설정 파일 로드 (이름 -> 설정 매핑)"""
        config_file = DEFAULT_CONFIG_PATH
        if config_file.exists():
            config = load_config(config_file, resolve_secrets=False)
            return {model['name']: model for model in config.get('models', [])}
        return {}

    def _get_model_price(self, model_name: str) -> Optional[Dict[str, float]]:
        """@description 모델 가격 정보 조회"""
        model_config = self._model_config.get(model_name, {})
        return model_config.get('price')

    def _load_token_usage(self) -> Dict:
        """@description 토큰 사용량 파일 로드"""
        token_file = self.problems_dir / self.token_usage_filename
        if token_file.exists():
            with open(token_file, 'r', encoding='utf-8') as f:
                return json.load(f)
        return {}

    def _get_token_usage(self, model_name: str, sheet_name: str) -> Optional[Dict[str, int]]:
        """@description 모델·시트별 토큰 사용량 조회"""
        models = self._token_usage.get('models', {})
        model_data = models.get(model_name, {})
        sections = model_data.get('sections', {})

        candidate_keys = [sheet_name]
        if sheet_name in ['영어', '한국사']:
            candidate_keys = [f'{sheet_name}-공통', sheet_name]
        elif sheet_name in ['물리1', '화학1', '생명1', '사회문화']:
            candidate_keys.append(f'탐구-{sheet_name}')

        for key in candidate_keys:
            if key not in sections:
                continue
            section_data = sections[key]
            return {
                'input_tokens': section_data.get('input_tokens', 0),
                'output_tokens': section_data.get('output_tokens', 0)
            }
        return None

    def export_to_json(self, sheet_name: str, model_name: str,
                       output_path: Path = None) -> Path:
        """@description Excel -> JSON 내보내기

        @param sheet_name 시트 이름
        @param model_name 모델 이름
        @param output_path 출력 경로(None 시 기본 경로)

        @return 저장 파일 경로
        """
        # 답안 추출
        answers = self.excel_handler.get_model_answers(sheet_name, model_name)

        # JSON 변환
        json_data = self.converter.excel_to_json(
            sheet_name, model_name, answers, self.excel_handler
        )

        # 출력 경로 결정
        if output_path is None:
            json_dir = self.path_mapper.sheet_to_json_path(sheet_name)
            if not json_dir:
                raise ValueError(f"'{sheet_name}'에 대한 경로 매핑이 없습니다.")
            output_path = json_dir / self.verified_filename

        # 출력 파일 존재 시 병합
        if output_path.exists():
            with open(output_path, 'r', encoding='utf-8') as f:
                existing_data = json.load(f)

            # 모델 점수·결과 병합
            json_model_name = self.model_mapper.excel_to_json(model_name)
            existing_data['model_scores'][json_model_name] = json_data['model_scores'][json_model_name]

            # results 같은 모델 결과 제거 후 추가
            existing_data['results'] = [
                r for r in existing_data['results']
                if r['model_name'] != json_model_name
            ]
            existing_data['results'].extend(json_data['results'])
            existing_data['timestamp'] = json_data['timestamp']

            json_data = existing_data

        # 저장
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(json_data, f, ensure_ascii=False, indent=2)

        print(f"내보내기 완료: {output_path}")
        return output_path

    def import_from_json(self, json_path: Path,
                         position: Optional[int] = None,
                         after_model: Optional[str] = None,
                         update_existing: bool = False,
                         excel_name: Optional[str] = None,
                         base_model: Optional[str] = None) -> bool:
        """@description JSON -> Excel 가져오기

        @param json_path results_verified.json 경로
        @param position 삽입 열 위치
        @param after_model 대상 모델 다음 삽입
        @param update_existing 기존 데이터 갱신 여부
        @param excel_name Excel 모델 이름(None 시 매핑 사용)
        @param base_model Excel 복사 기준 모델 이름

        @return 성공 여부
        """
        json_path = Path(json_path)
        if not json_path.exists():
            print(f"파일을 찾을 수 없습니다: {json_path}")
            return False

        # JSON 로드
        with open(json_path, 'r', encoding='utf-8') as f:
            json_data = json.load(f)

        # 시트 이름 결정
        sheet_name = self.path_mapper.json_to_sheet_name(json_path)
        if not sheet_name:
            print(f"경고: '{json_path}'에 대한 시트 매핑을 찾을 수 없습니다.")
            # subject/section 기준 시트 이름 추론
            subject = json_data.get('subject', '')
            section = json_data.get('section', '')
            if subject and section and subject != section:
                sheet_name = f"{subject}-{section}"
            elif subject:
                sheet_name = subject
            else:
                print("시트 이름을 결정할 수 없습니다.")
                return False

        # 시트 존재 확인
        if sheet_name not in self.excel_handler.get_sheet_names():
            print(f"시트를 찾을 수 없습니다: {sheet_name}")
            return False

        self.excel_handler.update_correct_answer_cells(
            sheet_name, self.converter.load_questions(sheet_name)
        )

        # 전체 모델 데이터 변환
        model_data_list = self.converter.json_to_excel(json_data)

        success_count = 0
        for model_name, answers, score in model_data_list:
            # Excel 모델 이름 결정(excel_name 지정 시 첫 모델 적용)
            if excel_name and success_count == 0:
                model_name = excel_name

            merged_answers = dict(answers)
            if base_model:
                existing_models = self.excel_handler.get_model_columns(sheet_name)
                if base_model not in existing_models:
                    raise ValueError(f"베이스 모델을 찾을 수 없습니다: {sheet_name} / {base_model}")
                base_answers = self.excel_handler.get_model_answers(sheet_name, base_model)
                merged_answers = dict(base_answers)
                merged_answers.update(answers)
                score = self.excel_handler.calculate_score_from_answers(sheet_name, merged_answers)

            # 대상 모델 존재 여부 확인
            existing_models = self.excel_handler.get_model_columns(sheet_name)

            if model_name in existing_models:
                if update_existing:
                    self.excel_handler.update_model_column(sheet_name, model_name, merged_answers, score)
                    print(f"업데이트 완료: {sheet_name} / {model_name} ({score}점)")
                    success_count += 1
                else:
                    print(f"'{model_name}' 모델이 이미 존재합니다. (건너뜀)")
            else:
                self.excel_handler.add_model_column(
                    sheet_name, model_name, merged_answers, score,
                    position=position, after_model=after_model
                )
                print(f"추가 완료: {sheet_name} / {model_name} ({score}점)")
                success_count += 1

        return success_count > 0

    def _infer_wrong_only_base_model(self, json_data: Dict, sheet_name: str) -> Optional[str]:
        """@description 오답 재평가 결과 파일 기반 Excel 기준 모델 추론"""
        model_scores = json_data.get('model_scores', {})
        if len(model_scores) != 1:
            return None

        model_name = next(iter(model_scores))
        base_model = self.WRONG_ONLY_BASE_MODELS.get(model_name)
        if not base_model:
            return None

        try:
            max_score = self.excel_handler.get_max_score(sheet_name)
        except Exception:
            return None

        total_points = json_data.get('total_points')
        total_verified = json_data.get('total_verified')
        if total_points is not None and total_points < max_score:
            return base_model

        # 점수 합 만점 충족 가능성 기준 문항 수 병행 확인
        try:
            questions = self.excel_handler._load_questions_for_sheet(sheet_name)
            expected_count = len(questions.get('questions', []))
        except Exception:
            expected_count = None
        if expected_count and total_verified is not None and total_verified < expected_count:
            return base_model

        return None

    def _materialize_missing_wrong_only_targets(self) -> int:
        """@description 오답 재평가 대상 모델 누락 시트 보완(기준 모델 답안 복사)

        오답 부재 과목: 오답 재평가 결과 JSON 생성 생략
        대상 모델 열 존재 시트 확인 후 나머지 시트 기준 모델 열 복사
        Excel·대시보드 전체 과목 표시용 답안 보완
        """
        materialized_count = 0
        sheets = self.path_mapper.get_all_sheets()
        sheet_models = {
            sheet_name: self.excel_handler.get_model_columns(sheet_name)
            for sheet_name in sheets
            if sheet_name in self.excel_handler.get_sheet_names()
        }

        for target_model, base_model in self.WRONG_ONLY_BASE_MODELS.items():
            target_exists_anywhere = any(
                target_model in models
                for models in sheet_models.values()
            )
            if not target_exists_anywhere:
                continue

            for sheet_name, models in sheet_models.items():
                if target_model in models:
                    continue
                if base_model not in models:
                    continue

                base_answers = self.excel_handler.get_model_answers(sheet_name, base_model)
                score = self.excel_handler.get_model_score(sheet_name, base_model)
                if score is None:
                    score = self.excel_handler.calculate_score_from_answers(sheet_name, base_answers)

                self.excel_handler.add_model_column(
                    sheet_name,
                    target_model,
                    dict(base_answers),
                    score,
                    after_model=base_model,
                )
                print(f"wrong-only 보정 추가: {sheet_name} / {target_model} <= {base_model} ({score}점)")
                materialized_count += 1

                # 다음 target 처리용 최신 열 상태 갱신
                sheet_models[sheet_name] = self.excel_handler.get_model_columns(sheet_name)

        return materialized_count

    def import_all(self, update_existing: bool = False) -> int:
        """@description 전체 검증 결과 JSON 가져오기"""
        count = 0
        for sheet_name in self.path_mapper.get_all_sheets():
            json_dir = self.path_mapper.sheet_to_json_path(sheet_name)
            if json_dir:
                json_file = json_dir / self.verified_filename
                if json_file.exists():
                    base_model = None
                    try:
                        with open(json_file, 'r', encoding='utf-8') as f:
                            json_data = json.load(f)
                        base_model = self._infer_wrong_only_base_model(json_data, sheet_name)
                        if base_model:
                            model_name = next(iter(json_data.get('model_scores', {})))
                            existing_models = self.excel_handler.get_model_columns(sheet_name)
                            if base_model not in existing_models:
                                print(
                                    f"경고: {sheet_name} / {model_name}은 부분 검증 파일이지만 "
                                    f"베이스 모델 '{base_model}' 컬럼이 없어 건너뜁니다."
                                )
                                continue
                            print(f"부분 검증 병합: {sheet_name} / {model_name} <= {base_model}")
                    except Exception as e:
                        print(f"경고: {json_file} 부분 검증 확인 실패 - {e}")

                    if self.import_from_json(
                        json_file,
                        update_existing=update_existing,
                        base_model=base_model,
                    ):
                        count += 1

        materialized_count = self._materialize_missing_wrong_only_targets()
        count += materialized_count

        if count > 0:
            self.excel_handler.save()

        return count

    def export_all_sheets_to_json(self, output_path: Path = None,
                                     model_name: str = None,
                                     all_models: bool = False) -> Path:
        """
        @description 전체 시트 단일 JSON 파일 내보내기(객체 배열)
        @param output_path 출력 경로(None 시 all_results.json)
        @param model_name 특정 모델 내보내기
        @param all_models 전체 모델 내보내기
        @return 저장 파일 경로
        """
        all_data = []

        for sheet_name in self.path_mapper.get_all_sheets():
            # 시트 존재 확인
            if sheet_name not in self.excel_handler.get_sheet_names():
                continue

            # questions.json 존재 확인
            json_dir = self.path_mapper.sheet_to_json_path(sheet_name)
            if not json_dir or not (json_dir / 'questions.json').exists():
                continue

            try:
                # 모델 목록 결정
                if model_name:
                    models_to_export = [model_name]
                elif all_models:
                    models_to_export = list(self.excel_handler.get_model_columns(sheet_name).keys())
                else:
                    models_to_export = list(self.excel_handler.get_model_columns(sheet_name).keys())

                for model in models_to_export:
                    try:
                        answers = self.excel_handler.get_model_answers(sheet_name, model)
                        json_data = self.converter.excel_to_json(
                            sheet_name, model, answers, self.excel_handler
                        )
                        # 공개 결과 형식 재구성
                        json_model_name = list(json_data['model_scores'].keys())[0]
                # results 불필요 필드 제거
                        clean_results = [
                            {
                                'question_number': r['question_number'],
                                'extracted_answer': r['extracted_answer'],
                                'correct_answer': r['correct_answer'],
                                'is_correct': r['is_correct'],
                                'points': r['points'],
                                'answer_status': r.get('answer_status'),
                                'provider_stop_reason': r.get('provider_stop_reason'),
                            }
                            for r in json_data['results']
                        ]
                        clean_data = {
                            'sheet_name': sheet_name,
                            'subject': json_data['subject'],
                            'section': json_data['section'],
                            'model_name': json_model_name,
                            'score': json_data['model_scores'][json_model_name],
                            'total_points': json_data['total_points'],
                            'correct_count': json_data['correct_count'],
                            'total_questions': json_data['total_verified'],
                        }
                        # 토큰 사용량 추가 (존재 시)
                        token_usage = self._get_token_usage(json_model_name, sheet_name)
                        if token_usage:
                            clean_data['token_usage'] = token_usage
                        # 가격 정보 추가 (존재 시)
                        price = self._get_model_price(json_model_name)
                        if price:
                            clean_data['price'] = price
                        clean_data['results'] = clean_results
                        all_data.append(clean_data)
                    except Exception as e:
                        print(f"경고: {sheet_name}/{model} 내보내기 실패 - {e}")

            except Exception as e:
                print(f"경고: {sheet_name} 처리 실패 - {e}")

        # 출력 경로 결정
        if output_path is None:
            output_path = Path(self.all_results_filename)

        # 저장
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(all_data, f, ensure_ascii=False, indent=2)

        if _sync_model_metadata(self._model_config, self.model_metadata_path):
            print(f"모델 메타데이터 동기화 완료: {self.model_metadata_path}")

        print(f"내보내기 완료: {output_path} ({len(all_data)}개 항목)")
        return output_path

    def export_hard_all_sections_to_json(self, output_path: Path = None,
                                         model_name: str = None) -> Path:
        """
        @description hard Excel 전체 모델 열 → hard_all_results.json 형식 내보내기
        hard_results*.json 개별 실행 원본 → 전체 집계 기준 제외
        대시보드 전체 집계: hard Excel 전체 모델 기준
        """
        return self.export_all_sheets_to_json(
            output_path=output_path,
            model_name=model_name,
            all_models=model_name is None,
        )

    def list_models(self, sheet_name: str = None) -> Dict[str, List[str]]:
        """@description 모델 목록 반환"""
        result = {}

        if sheet_name:
            sheets = [sheet_name]
        else:
            sheets = self.excel_handler.get_sheet_names()

        for sheet in sheets:
            if sheet in self.path_mapper.SHEET_TO_JSON or sheet_name:
                try:
                    models = list(self.excel_handler.get_model_columns(sheet).keys())
                    result[sheet] = models
                except:
                    pass

        return result

    def validate(self, sheet_name: str = None) -> List[str]:
        """@description 데이터 일관성 검증"""
        issues = []

        if sheet_name:
            sheets = [sheet_name]
        else:
            sheets = self.path_mapper.get_all_sheets()

        for sheet in sheets:
            json_dir = self.path_mapper.sheet_to_json_path(sheet)
            if not json_dir:
                continue

            # questions.json 확인
            questions_file = json_dir / 'questions.json'
            if not questions_file.exists():
                issues.append(f"[{sheet}] questions.json 없음: {questions_file}")
                continue

            with open(questions_file, 'r', encoding='utf-8') as f:
                questions = json.load(f)

            expected_count = len(questions['questions'])
            expected_total = sum(q['points'] for q in questions['questions'])

            # Excel 데이터 확인
            try:
                model_columns = self.excel_handler.get_model_columns(sheet)
                max_score = self.excel_handler.get_max_score(sheet)

                if max_score != expected_total:
                    issues.append(f"[{sheet}] 만점 불일치: Excel={max_score}, JSON={expected_total}")

                for model_name in model_columns:
                    answers = self.excel_handler.get_model_answers(sheet, model_name)
                    if len(answers) != expected_count:
                        issues.append(
                            f"[{sheet}] {model_name}: 문항 수 불일치 "
                            f"(Excel={len(answers)}, JSON={expected_count})"
                        )
            except Exception as e:
                issues.append(f"[{sheet}] 검증 오류: {e}")

        return issues


def main():
    """@description 시험·쉬움 모드 동기화와 모델 메타데이터 명령 실행"""
    arguments = sys.argv[1:]
    if arguments and arguments[0] == "metadata":
        parser = argparse.ArgumentParser(
            description="모델 공개 메타데이터 동기화", allow_abbrev=False
        )
        parser.add_argument("metadata")
        parser.add_argument("--config", required=True, help="모델 설정 파일 경로")
        parser.add_argument("--output", help="모델 메타데이터 출력 경로")
        args = parser.parse_args(arguments)
        try:
            config = load_config(args.config, resolve_secrets=False)
            models = {model["name"]: model for model in config.get("models", [])}
            output_path = (
                Path(args.output).expanduser().resolve()
                if args.output
                else _metadata_path_for_config(args.config)
            )
            changed = _sync_model_metadata(models, output_path)
            state = "변경" if changed else "변경 없음"
            print(f"모델 메타데이터 동기화 완료 ({state}): {output_path}")
            return 0
        except (OSError, ValueError) as error:
            print(f"오류: {error}", file=sys.stderr)
            return 2
    return _generic_sync_main(arguments)

    arguments = sys.argv[1:]
    if any(argument == '--exam' or argument.startswith('--exam=') for argument in arguments):
        return _generic_sync_main(arguments)
    parser = argparse.ArgumentParser(
        description='Excel-JSON 양방향 동기화 도구',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
예시:
  # JSON -> Excel 가져오기
  python sync_data.py import --json problems/국어/공통/results_verified.json
  python sync_data.py import --all
  python sync_data.py import --all --hard

  # Excel -> JSON 내보내기 (단일 시트)
  python sync_data.py export --sheet 국어-공통 --model "GPT-5.1"
  python sync_data.py export --sheet 국어-공통 --all-models

  # Excel -> JSON 내보내기 (모든 시트를 하나의 JSON 배열로)
  python sync_data.py export --all-sheets
  python sync_data.py export --all-sheets --output all_results.json
  python sync_data.py export --all-sheets --model "GPT-5.1"  # 특정 모델만
  python sync_data.py export --all-sheets --hard

  # 모델 목록 확인
  python sync_data.py list
  python sync_data.py list --sheet 국어-공통

  # 검증
  python sync_data.py validate
        """
    )

    subparsers = parser.add_subparsers(dest='command', help='명령')

    # Export 명령
    export_parser = subparsers.add_parser('export', help='Excel -> JSON 내보내기')
    export_parser.add_argument('--sheet', help='시트 이름 (예: 국어-공통)')
    export_parser.add_argument('--all-sheets', action='store_true', help='모든 시트를 하나의 JSON 배열로 내보내기')
    export_parser.add_argument('--model', help='모델 이름')
    export_parser.add_argument('--all-models', action='store_true', help='모든 모델 내보내기')
    export_parser.add_argument('--output', help='출력 파일 경로')
    export_parser.add_argument('--hard', action='store_true', help='hard 전용 Excel/JSON 파일 사용')

    # Import 명령
    import_parser = subparsers.add_parser('import', help='JSON -> Excel 가져오기')
    import_parser.add_argument('--json', help='results_verified.json 경로')
    import_parser.add_argument('--all', action='store_true', help='모든 JSON 파일 가져오기')
    import_parser.add_argument('--position', type=int, help='삽입할 열 위치')
    import_parser.add_argument('--after', help='이 모델 다음에 삽입')
    import_parser.add_argument('--update', action='store_true', help='기존 데이터 업데이트')
    import_parser.add_argument('--excel-name', help='Excel에서 사용할 모델 이름')
    import_parser.add_argument('--base-model', help='기존 Excel 컬럼을 베이스로 복사한 뒤 JSON 답안만 덮어쓰기')
    import_parser.add_argument('--hard', action='store_true', help='hard 전용 Excel/JSON 파일 사용')

    # List 명령
    list_parser = subparsers.add_parser('list', help='모델 목록 확인')
    list_parser.add_argument('--sheet', help='특정 시트만')
    list_parser.add_argument('--hard', action='store_true', help='hard 전용 Excel 사용')

    # Validate 명령
    validate_parser = subparsers.add_parser('validate', help='데이터 검증')
    validate_parser.add_argument('--sheet', help='특정 시트만')
    validate_parser.add_argument('--hard', action='store_true', help='hard 전용 Excel/JSON 파일 사용')

    metadata_parser = subparsers.add_parser('metadata', help='모델 공개 메타데이터 동기화')
    metadata_parser.add_argument('--config', required=True, help='모델 설정 파일 경로')
    metadata_parser.add_argument('--output', help='모델 메타데이터 출력 경로')

    # 공통 옵션
    parser.add_argument('--excel', default=None,
                        help='Excel 파일 경로')
    parser.add_argument('--mapping', default='model_mapping.json',
                        help='모델 이름 매핑 파일')

    args = parser.parse_args()

    if not args.command:
        parser.print_help()
        return

    if args.command == 'metadata':
        try:
            config = load_config(args.config, resolve_secrets=False)
            models = {model['name']: model for model in config.get('models', [])}
            output_path = (
                Path(args.output).expanduser().resolve()
                if args.output
                else _metadata_path_for_config(args.config)
            )
            changed = _sync_model_metadata(models, output_path)
            state = '변경' if changed else '변경 없음'
            print(f'모델 메타데이터 동기화 완료 ({state}): {output_path}')
            return 0
        except (OSError, ValueError) as error:
            print(f'오류: {error}', file=sys.stderr)
            return 2

    hard_mode = getattr(args, 'hard', False)
    excel_path = Path(args.excel) if args.excel else (
        DEFAULT_HARD_EXCEL_PATH if hard_mode else DEFAULT_EXCEL_PATH
    )
    if hard_mode and not excel_path.exists():
        _create_hard_excel_template(DEFAULT_EXCEL_PATH, excel_path)

    # SyncManager 초기화
    sync = SyncManager(
        excel_path=excel_path,
        model_mapping_path=Path(args.mapping),
        hard_mode=hard_mode
    )

    # 명령 실행
    if args.command == 'export':
        if args.all_sheets:
                # 전체 시트 단일 JSON 배열 내보내기
            output = Path(args.output) if args.output else None
            if hard_mode:
                sync.export_hard_all_sections_to_json(
                    output_path=output,
                    model_name=args.model,
                )
            else:
                sync.export_all_sheets_to_json(
                    output_path=output,
                    model_name=args.model,
                    all_models=args.all_models or (args.model is None)
                )
        elif args.sheet:
            if args.all_models:
                models = sync.list_models(args.sheet).get(args.sheet, [])
                for model in models:
                    sync.export_to_json(args.sheet, model)
            elif args.model:
                output = Path(args.output) if args.output else None
                sync.export_to_json(args.sheet, args.model, output)
            else:
                print("--model 또는 --all-models 옵션을 지정하세요.")
        else:
            print("--sheet 또는 --all-sheets 옵션을 지정하세요.")

    elif args.command == 'import':
        if args.all:
            count = sync.import_all(update_existing=args.update)
            print(f"\n총 {count}개 파일 가져오기 완료")
        elif args.json:
            success = sync.import_from_json(
                Path(args.json),
                position=args.position,
                after_model=args.after,
                update_existing=args.update,
                excel_name=args.excel_name,
                base_model=args.base_model
            )
            if success:
                sync.excel_handler.save()
        else:
            print("--json 또는 --all 옵션을 지정하세요.")

    elif args.command == 'list':
        models = sync.list_models(args.sheet)
        for sheet, model_list in models.items():
            print(f"\n[{sheet}] ({len(model_list)}개 모델)")
            for model in model_list:
                score = sync.excel_handler.get_model_score(sheet, model)
                print(f"  - {model}: {score}점")

    elif args.command == 'validate':
        issues = sync.validate(args.sheet)
        if issues:
            print("검증 결과: 문제 발견")
            for issue in issues:
                print(f"  - {issue}")
        else:
            print("검증 결과: 정상")


if __name__ == '__main__':
    raise SystemExit(main())
