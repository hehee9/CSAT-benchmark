"""@description 비공개 검증 결과·Excel·공개 결과 사이의 단일 실행 변환"""

from __future__ import annotations

import copy
import json
import math
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook, load_workbook

from .answers import CorrectAnswer, _is_correct_answer
from .configuration import load_config
from .evaluation import load_verified, save_verified
from .exams import ExamManifest, SectionManifest, load_exam, load_section_questions
from .metadata import merge_model_metadata
from .runner import RunnerError, select_sections
from .runs import load_run, result_identity, result_is_completed


_EXCEL_MAPPING_FILE = "model_mapping.json"
_ANSWER_HEADERS = {"문항 번호", "정답", "배점"}
_NON_MODEL_HEADERS = _ANSWER_HEADERS | {"총점", "총합", "점수"}
_REFUSAL_MARKERS = {"-2", "(검열)", "검열", "refusal", "Refusal"}
_NO_ANSWER_MARKERS = {"-1", "포기", "(포기)"}


def _format_excel_correct_answer(value: CorrectAnswer) -> int | str:
    """@description 정답 단일값·복수값을 Excel 표시값으로 변환"""
    if isinstance(value, list):
        return ", ".join(str(answer) for answer in value)
    return value


def target_to_section_key(
    target: str,
    subject: str | None = None,
    section: str | None = None,
) -> str:
    """@description 정식 target을 기존 Excel·공개 섹션 키로 변환"""
    if target in {"영어/전체", "한국사/전체"}:
        return target.split("/", 1)[0]
    if target.startswith("탐구/"):
        return target.split("/", 1)[1]
    if target == "수학/미적분":
        return "수학-미적"
    return target.replace("/", "-")


def _token_section_key(target: str, input_mode: str) -> str:
    """@description 실행 방식에 맞는 기존 토큰 섹션 키 반환"""
    if target in {"영어/전체", "한국사/전체"} and input_mode == "section":
        return target.split("/", 1)[0] + "-공통"
    return target_to_section_key(target)


def _load_json(path: Path, default: Any) -> Any:
    """@description JSON 파일 로드 및 파일 부재 시 기본값 반환"""
    if not path.is_file():
        return copy.deepcopy(default)
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, payload: Any) -> Path:
    """@description UTF-8 JSON 파일 저장"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def _write_json_if_changed(path: Path, payload: Any) -> Path:
    """@description JSON 의미가 달라질 때만 파일 교체"""
    if path.is_file():
        if json.loads(path.read_text(encoding="utf-8")) == payload:
            return path
    return _write_json(path, payload)


def _as_exam(exam: ExamManifest | str | Path | None) -> ExamManifest:
    """@description 시험 입력을 ExamManifest로 변환"""
    if exam is None:
        raise ValueError("Excel 변환에는 exam 매니페스트가 필요합니다.")
    return exam if isinstance(exam, ExamManifest) else load_exam(exam)


def _as_run(run: Mapping[str, Any] | str | Path | None) -> dict[str, Any] | None:
    """@description 실행 입력을 복사된 정본 문서로 변환"""
    if run is None:
        return None
    return load_run(run) if isinstance(run, (str, Path)) else copy.deepcopy(dict(run))


def _selected_scope(
    manifest: ExamManifest,
    run: Mapping[str, Any],
    *,
    model_names: Sequence[str] | None,
    targets: Sequence[str] | None,
    subject: str | None,
    section: str | None,
    subjects: Sequence[str] | None,
    benchmark_all: bool,
    question_numbers: Sequence[int] | None,
) -> tuple[list[SectionManifest], list[str]]:
    """@description 실행에 저장된 범위에서 동기화 선택 범위 결정"""
    stored_models = [model["name"] for model in run["selected_models"]]
    if model_names is None:
        selected_models = stored_models
    else:
        requested = list(dict.fromkeys(model_names))
        unknown = set(requested) - set(stored_models)
        if unknown:
            raise ValueError(f"실행에 없는 모델입니다: {', '.join(sorted(unknown))}")
        selected_models = [name for name in stored_models if name in requested]
    try:
        sections = select_sections(
            manifest,
            targets=targets,
            subject=subject,
            section=section,
            subjects=subjects,
            benchmark_all=benchmark_all,
        )
    except RunnerError as error:
        raise ValueError(str(error)) from error
    stored_targets = set(run["selected_targets"])
    if any(item.target not in stored_targets for item in sections):
        raise ValueError("선택한 섹션이 실행에 저장되어 있지 않습니다.")
    if question_numbers is not None:
        if run["input_mode"] != "question":
            raise ValueError("--question-numbers는 문항별 실행에서만 사용할 수 있습니다.")
        if any(
            isinstance(number, bool) or not isinstance(number, int) or number < 1
            for number in question_numbers
        ):
            raise ValueError("문항 번호는 양의 정수여야 합니다.")
    return sections, selected_models


def _stored_question_numbers(run: Mapping[str, Any], target: str) -> list[int] | None:
    """@description 실행 문맥에 기록된 target별 문항 번호 반환"""
    for context in run.get("input_context", []):
        if isinstance(context, Mapping) and context.get("target") == target:
            values = context.get("question_numbers")
            if values is None:
                return None
            if (
                not isinstance(values, list)
                or any(
                    isinstance(value, bool) or not isinstance(value, int) or value < 1
                    for value in values
                )
                or len(values) != len(set(values))
            ):
                raise ValueError(f"실행 input_context의 문항 번호가 잘못되었습니다: {target}")
            return list(values)
    return None


def _section_aliases(section: SectionManifest) -> set[str]:
    """@description 매니페스트 섹션에 허용된 기존 이름 목록 생성"""
    aliases = {
        section.target,
        f"{section.subject}/{section.section}",
        f"{section.subject}-{section.section}",
        section.target.replace("/", "-"),
        target_to_section_key(section.target, section.subject, section.section),
    }
    if section.kind == "subject":
        aliases.update({f"{section.subject}/{section.subject}", section.subject})
    if section.group == "탐구":
        aliases.update({f"탐구/{section.subject}", f"탐구-{section.subject}", section.subject})
    if section.subject == "수학" and section.section in {"미적", "미적분"}:
        aliases.update({"수학/미적", "수학/미적분", "수학-미적", "수학-미적분"})
    return aliases


def _record_target(record: Mapping[str, Any], manifest: ExamManifest) -> str | None:
    """@description 과거 공개 행의 target 별칭을 매니페스트 target으로 정규화"""
    values = [record.get("target"), record.get("sheet_name")]
    subject = record.get("subject")
    section_name = record.get("section")
    if isinstance(subject, str) and isinstance(section_name, str):
        values.extend((f"{subject}/{section_name}", f"{subject}-{section_name}"))
    for value in values:
        if not isinstance(value, str) or not value:
            continue
        for section in manifest.sections:
            if value in _section_aliases(section):
                return section.target
    if isinstance(subject, str):
        matches = [
            section
            for section in manifest.sections
            if section.kind == "subject" and section.subject == subject
        ]
        if len(matches) == 1:
            return matches[0].target
        inquiry = [
            section
            for section in manifest.sections
            if section.group == "탐구" and section.subject == subject
        ]
        if len(inquiry) == 1:
            return inquiry[0].target
    return None


def _excel_sheet_for_section(section: SectionManifest) -> str:
    """@description 매니페스트 섹션의 기존 Excel 시트명 반환"""
    return target_to_section_key(section.target, section.subject, section.section)


def _metadata_section_key(section: SectionManifest) -> str:
    """@description 기존 문항 메타데이터의 섹션 키 반환"""
    return f"{section.subject}-{section.section}"


def _find_header_row(worksheet) -> int:
    """@description Excel 시트의 동적 헤더 행 조회"""
    for row_index in range(1, min(20, worksheet.max_row) + 1):
        for cell in worksheet[row_index]:
            if str(cell.value).strip() == "문항 번호":
                return row_index
    raise ValueError(f"'{worksheet.title}' 시트에서 헤더 행을 찾을 수 없습니다.")


def _find_score_row(worksheet, header_row: int) -> int:
    """@description Excel 시트의 총점 행 조회"""
    for row_index in range(header_row + 1, worksheet.max_row + 1):
        value = worksheet.cell(row_index, column=1).value
        if str(value).strip() in {"총점", "총합", "점수"}:
            return row_index
    return worksheet.max_row + 1


def _load_model_mapping(root: Path) -> dict[str, str]:
    """@description JSON 모델명·Excel 열 이름 매핑 로드"""
    path = root / _EXCEL_MAPPING_FILE
    if not path.is_file():
        return {}
    payload = _load_json(path, {})
    if not isinstance(payload, Mapping):
        raise ValueError(f"모델 매핑 파일은 객체여야 합니다: {path}")
    return {str(key): str(value) for key, value in payload.items()}


def _sheet_target(workbook, sheet_name: str, manifest: ExamManifest | None = None) -> str | None:
    """@description 숨김 메타데이터·매니페스트 기준 Excel 시트 target 조회"""
    if "_csat_meta" in workbook.sheetnames:
        meta = workbook["_csat_meta"]
        for row in meta.iter_rows(min_row=2, values_only=True):
            if row and row[0] == sheet_name and isinstance(row[1], str):
                return row[1]
    if manifest is not None:
        matches = [
            section.target
            for section in manifest.sections
            if sheet_name in _section_aliases(section)
        ]
        if len(matches) == 1:
            return matches[0]
    return None


def _worksheet_for_target(workbook, manifest: ExamManifest, target: str):
    """@description manifest target에 대응하는 Excel 시트 조회"""
    for sheet_name in workbook.sheetnames:
        if sheet_name == "_csat_meta":
            continue
        if _sheet_target(workbook, sheet_name, manifest) == target:
            return workbook[sheet_name]
    return None


def _ensure_meta(workbook):
    """@description 새 Excel 통합 문서의 target 메타데이터 시트 확보"""
    if "_csat_meta" in workbook.sheetnames:
        return workbook["_csat_meta"]
    meta = workbook.create_sheet("_csat_meta")
    meta.sheet_state = "hidden"
    meta.append(["sheet_name", "target"])
    return meta


def _new_sheet(
    workbook,
    section: SectionManifest,
    models: Sequence[str],
    mapping: Mapping[str, str],
    manifest: ExamManifest,
):
    """@description 새 Excel 섹션 시트 생성"""
    sheet_name = _excel_sheet_for_section(section)[:31] or "결과"
    original = sheet_name
    suffix = 1
    while sheet_name in workbook.sheetnames:
        suffix += 1
        sheet_name = f"{original[:27]}-{suffix}"
    meta = _ensure_meta(workbook)
    meta.append([sheet_name, section.target])
    worksheet = workbook.create_sheet(sheet_name)
    headers = ["문항 번호", "정답", "배점"]
    headers.extend(mapping.get(model, model) for model in models)
    worksheet.append(headers)
    questions = load_section_questions(manifest, section)
    for question in questions:
        worksheet.append(
            [question.number, _format_excel_correct_answer(question.correct_answer), question.points]
            + [None] * len(models)
        )
    worksheet.append(["총점", None, sum(question.points for question in questions)])
    return worksheet


def _answers_by_model(
    verified_data: Mapping[str, Any],
    sections: Sequence[SectionManifest],
    models: Sequence[str],
    question_numbers: Sequence[int] | None,
) -> dict[tuple[str, str], dict[int, Mapping[str, Any]]]:
    """@description verified flat 행을 target·모델·문항별로 그룹화"""
    allowed_targets = {section.target for section in sections}
    allowed_numbers = set(question_numbers) if question_numbers is not None else None
    grouped: dict[tuple[str, str], dict[int, Mapping[str, Any]]] = {}
    for row in verified_data.get("results", []):
        if not isinstance(row, Mapping):
            continue
        target = row.get("target")
        model = row.get("model_name")
        number = row.get("question_number")
        if (
            target in allowed_targets
            and model in models
            and isinstance(number, int)
            and number > 0
            and (allowed_numbers is None or number in allowed_numbers)
        ):
            grouped.setdefault((target, model), {})[number] = row
    return grouped


def export_run_to_excel(
    verified: Mapping[str, Any] | str | Path,
    excel_path: str | Path,
    *,
    run: Mapping[str, Any] | str | Path | None = None,
    exam: ExamManifest | str | Path | None = None,
    model_names: Sequence[str] | None = None,
    targets: Sequence[str] | None = None,
    subject: str | None = None,
    section: str | None = None,
    subjects: Sequence[str] | None = None,
    benchmark_all: bool = False,
    question_numbers: Sequence[int] | None = None,
) -> Path:
    """@description flat verified 답안을 기존 또는 새 Excel의 모델별 한 열에 기록"""
    manifest = _as_exam(exam)
    verified_data = load_verified(verified) if isinstance(verified, (str, Path)) else copy.deepcopy(dict(verified))
    run_data = _as_run(run)
    if run_data is None:
        run_data = {
            "selected_models": [{"name": name} for name in verified_data.get("selected_models", [])],
            "selected_targets": list(verified_data.get("selected_targets", [])),
            "input_mode": "question",
        }
    sections, models = _selected_scope(
        manifest,
        run_data,
        model_names=model_names,
        targets=targets,
        subject=subject,
        section=section,
        subjects=subjects,
        benchmark_all=benchmark_all,
        question_numbers=question_numbers,
    )
    output_path = Path(excel_path).expanduser().resolve()
    workbook_exists = output_path.is_file()
    workbook = load_workbook(output_path) if workbook_exists else Workbook()
    if not workbook_exists and workbook.active is not None:
        workbook.remove(workbook.active)
    mapping = _load_model_mapping(manifest.project_root)
    grouped = _answers_by_model(verified_data, sections, models, question_numbers)
    for section_manifest in sections:
        worksheet = _worksheet_for_target(workbook, manifest, section_manifest.target)
        if worksheet is None:
            worksheet = _new_sheet(workbook, section_manifest, models, mapping, manifest)
        header_row = _find_header_row(worksheet)
        score_row = _find_score_row(worksheet, header_row)
        headers = [worksheet.cell(header_row, column).value for column in range(1, worksheet.max_column + 1)]
        header_to_column = {
            str(value).strip(): index
            for index, value in enumerate(headers, start=1)
            if value is not None and str(value).strip()
        }
        answer_column = header_to_column.get("정답", 2)
        expected_questions = load_section_questions(manifest, section_manifest)
        question_rows = {
            int(worksheet.cell(row, 1).value): row
            for row in range(header_row + 1, min(score_row, worksheet.max_row + 1))
            if isinstance(worksheet.cell(row, 1).value, int)
            and not isinstance(worksheet.cell(row, 1).value, bool)
        }
        for question in expected_questions:
            if question.number not in question_rows:
                row_index = score_row
                worksheet.insert_rows(row_index)
                worksheet.cell(row_index, 1).value = question.number
                worksheet.cell(row_index, answer_column).value = _format_excel_correct_answer(
                    question.correct_answer
                )
                if "배점" in header_to_column:
                    worksheet.cell(row_index, header_to_column["배점"]).value = question.points
                score_row += 1
                question_rows[question.number] = row_index
            worksheet.cell(
                question_rows[question.number], answer_column
            ).value = _format_excel_correct_answer(question.correct_answer)
        for model in models:
            excel_model = mapping.get(model, model)
            column = header_to_column.get(excel_model) or header_to_column.get(model)
            if column is None:
                column = worksheet.max_column + 1
                worksheet.cell(header_row, column).value = excel_model
                header_to_column[excel_model] = column
            for question in expected_questions:
                row_index = question_rows[question.number]
                row = grouped.get((section_manifest.target, model), {}).get(question.number)
                if row is not None and row.get("extracted_answer") is not None:
                    cell = worksheet.cell(row_index, column)
                    answer = row.get("extracted_answer")
                    if answer == -1:
                        cell.value = "(포기)"
                        cell.number_format = "General"
                    elif answer == -2:
                        cell.value = "(검열)"
                        cell.number_format = "General"
                    else:
                        cell.value = answer
            score_cell = worksheet.cell(score_row, column)
            if score_cell.data_type != "f":
                score_cell.value = sum(
                    int(answer_row["points"])
                    for answer_row in grouped.get((section_manifest.target, model), {}).values()
                    if answer_row.get("is_correct") is True
                )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def _normalise_excel_answer(value: Any) -> tuple[int | None, str | None]:
    """@description Excel 셀 값을 답안 숫자·상태로 변환"""
    if value is None or str(value).strip() == "":
        return None, None
    text = str(value).strip()
    if text in _NO_ANSWER_MARKERS:
        return -1, "no_answer"
    if text in _REFUSAL_MARKERS:
        return -2, "refusal"
    try:
        return int(value), "answered"
    except (TypeError, ValueError):
        return None, None


def _model_name_from_header(header: str, models: Sequence[str], mapping: Mapping[str, str]) -> str | None:
    """@description Excel 모델 열 이름을 정본 모델명으로 변환"""
    reverse = {excel: model for model, excel in mapping.items()}
    if header in models:
        return header
    model = reverse.get(header)
    return model if model in models else None


def import_excel_corrections(
    verified: Mapping[str, Any] | str | Path,
    excel_path: str | Path,
    *,
    output_path: str | Path | None = None,
    exam: ExamManifest | str | Path | None = None,
    run: Mapping[str, Any] | str | Path | None = None,
    model_names: Sequence[str] | None = None,
    targets: Sequence[str] | None = None,
    subject: str | None = None,
    section: str | None = None,
    subjects: Sequence[str] | None = None,
    benchmark_all: bool = False,
    question_numbers: Sequence[int] | None = None,
) -> dict[str, Any]:
    """@description Excel 수동 답안 변경을 flat verified 행에 재채점하여 저장"""
    manifest = _as_exam(exam)
    verified_path = Path(verified).expanduser().resolve() if isinstance(verified, (str, Path)) else None
    verified_data = load_verified(verified) if verified_path is not None else copy.deepcopy(dict(verified))
    run_data = _as_run(run)
    if run_data is None:
        run_data = {
            "selected_models": [{"name": name} for name in verified_data.get("selected_models", [])],
            "selected_targets": list(verified_data.get("selected_targets", [])),
            "input_mode": "question",
        }
    sections, models = _selected_scope(
        manifest,
        run_data,
        model_names=model_names,
        targets=targets,
        subject=subject,
        section=section,
        subjects=subjects,
        benchmark_all=benchmark_all,
        question_numbers=question_numbers,
    )
    workbook = load_workbook(Path(excel_path).expanduser().resolve(), data_only=True)
    mapping = _load_model_mapping(manifest.project_root)
    rows_by_key = {
        result_identity(row): copy.deepcopy(dict(row))
        for row in verified_data.get("results", [])
        if isinstance(row, Mapping)
    }
    changed = 0
    selected_targets = {item.target for item in sections}
    requested_numbers = set(question_numbers) if question_numbers is not None else None
    source_correct_answers = {
        section_manifest.target: {
            question.number: question.correct_answer
            for question in load_section_questions(manifest, section_manifest)
        }
        for section_manifest in sections
    }
    for worksheet in workbook.worksheets:
        if worksheet.title == "_csat_meta":
            continue
        target = _sheet_target(workbook, worksheet.title, manifest)
        if target not in selected_targets:
            continue
        header_row = _find_header_row(worksheet)
        score_row = _find_score_row(worksheet, header_row)
        headers = [worksheet.cell(header_row, column).value for column in range(1, worksheet.max_column + 1)]
        number_column = next(
            (index for index, value in enumerate(headers, start=1) if str(value).strip() == "문항 번호"),
            1,
        )
        model_columns = [
            (index, _model_name_from_header(str(value).strip(), models, mapping))
            for index, value in enumerate(headers, start=1)
            if value is not None
            and str(value).strip() not in _NON_MODEL_HEADERS
            and not str(value).strip().startswith("Unnamed")
        ]
        for row_index in range(header_row + 1, min(score_row, worksheet.max_row + 1)):
            number = worksheet.cell(row_index, number_column).value
            if isinstance(number, bool) or not isinstance(number, int) or number < 1:
                continue
            if requested_numbers is not None and number not in requested_numbers:
                continue
            for column, model_name in model_columns:
                if model_name is None:
                    continue
                item = rows_by_key.get((target, model_name, number))
                if item is None or item.get("complete") is not True:
                    continue
                correct_answer = source_correct_answers[target].get(number)
                if correct_answer is None:
                    continue
                answer, status = _normalise_excel_answer(worksheet.cell(row_index, column).value)
                if answer is None or status is None:
                    continue
                if (
                    item.get("extracted_answer") == answer
                    and item.get("correct_answer") == correct_answer
                ):
                    continue
                item.update(
                    {
                        "extracted_answer": answer,
                        "correct_answer": copy.deepcopy(correct_answer),
                        "is_correct": _is_correct_answer(answer, correct_answer),
                        "answer_status": status,
                        "provenance": "excel_manual",
                        "needs_manual_review": False,
                    }
                )
                changed += 1
    verified_data["results"] = [rows_by_key[key] for key in sorted(rows_by_key)]
    score_by_model = {name: 0 for name in verified_data.get("selected_models", models)}
    complete_by_model = {name: True for name in score_by_model}
    for row in verified_data["results"]:
        model = row.get("model_name")
        score_by_model.setdefault(model, 0)
        complete_by_model.setdefault(model, True)
        if row.get("complete") is True and row.get("is_correct") is True:
            score_by_model[model] += int(row.get("points", 0))
        if row.get("complete") is not True:
            complete_by_model[model] = False
    verified_data["score_by_model"] = score_by_model
    verified_data["complete_by_model"] = complete_by_model
    destination = (
        Path(output_path).expanduser().resolve()
        if output_path is not None
        else verified_path
    )
    if destination is None:
        raise ValueError("매핑 입력에는 output_path가 필요합니다.")
    save_verified(verified_data, destination)
    return {"verified": verified_data, "path": destination, "changed": changed}


def _numeric(value: Any) -> bool:
    """@description 유한한 토큰 수치 여부 확인"""
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def _new_token_record(rows: Iterable[Mapping[str, Any]]) -> dict[str, int | str | None]:
    """@description 단일 실행 원본 결과에서 토큰 합계 생성"""
    rows = list(rows)
    output: dict[str, int | str | None] = {
        "input_tokens": None,
        "output_tokens": None,
        "total_tokens": None,
    }
    for source in ("input_tokens", "output_tokens", "total_tokens"):
        values = [row.get(source) for row in rows]
        if values and all(_numeric(value) for value in values):
            output[source] = int(sum(values))
    output["question_count"] = len(rows)
    timestamps = [
        row["timestamp"]
        for row in rows
        if isinstance(row.get("timestamp"), str) and row["timestamp"]
    ]
    if timestamps:
        output["last_updated"] = max(timestamps)
    return output


def _merge_token_scope(
    token_data: Mapping[str, Any],
    updates: Mapping[tuple[str, str], Mapping[str, Any]],
) -> dict[str, Any]:
    """@description 기존 토큰 기록에 변경된 target만 단일 실행 합계로 반영"""
    merged = copy.deepcopy(dict(token_data))
    models = merged.setdefault("models", {})
    if not isinstance(models, dict):
        raise ValueError("기존 token_usage.json의 models는 객체여야 합니다.")
    for (model_name, section_key), record in updates.items():
        model = copy.deepcopy(models.get(model_name, {}))
        if not isinstance(model, dict):
            model = {}
        sections = model.setdefault("sections", {})
        if not isinstance(sections, dict):
            sections = {}
            model["sections"] = sections
        section_record = dict(record)
        if "last_updated" not in section_record:
            existing_section = sections.get(section_key)
            if isinstance(existing_section, Mapping) and isinstance(
                existing_section.get("last_updated"), str
            ):
                section_record["last_updated"] = existing_section["last_updated"]
        sections[section_key] = section_record
        section_records = list(sections.values())
        for source, destination in (
            ("input_tokens", "total_input_tokens"),
            ("output_tokens", "total_output_tokens"),
            ("total_tokens", "total_tokens"),
        ):
            values = [
                item.get(source) if isinstance(item, Mapping) else None
                for item in section_records
            ]
            if values and all(_numeric(value) for value in values):
                model[destination] = int(sum(values))
            else:
                model[destination] = None
        counts = [item.get("question_count") for item in sections.values() if isinstance(item, Mapping)]
        if counts and all(_numeric(value) for value in counts):
            model["question_count"] = int(sum(counts))
        timestamps = [
            item["last_updated"]
            for item in section_records
            if isinstance(item, Mapping)
            and isinstance(item.get("last_updated"), str)
            and item["last_updated"]
        ]
        if timestamps:
            model["last_updated"] = max(timestamps)
        models[model_name] = model
    return merged


def _public_record(
    rows: Sequence[Mapping[str, Any]],
    section: SectionManifest,
    model_snapshot: Mapping[str, Any] | None,
) -> dict[str, Any]:
    """@description 완료된 flat verified 행을 기존 공개 레코드로 변환"""
    ordered = sorted(rows, key=lambda row: int(row["question_number"]))
    record = {
        "sheet_name": _excel_sheet_for_section(section),
        "target": section.target,
        "subject": section.subject,
        "section": section.section,
        "model_name": ordered[0]["model_name"],
        "score": sum(int(row["points"]) for row in ordered if row.get("is_correct") is True),
        "total_points": sum(int(row["points"]) for row in ordered),
        "correct_count": sum(1 for row in ordered if row.get("is_correct") is True),
        "total_questions": len(ordered),
        "complete": True,
        "results": [
            {
                key: row.get(key)
                for key in (
                    "question_number",
                    "extracted_answer",
                    "correct_answer",
                    "is_correct",
                    "points",
                    "answer_status",
                    "provider_stop_reason",
                )
            }
            for row in ordered
        ],
    }
    if model_snapshot and model_snapshot.get("price") is not None:
        record["price"] = copy.deepcopy(model_snapshot["price"])
    return record


def publish_run(
    exam: ExamManifest | str | Path,
    run: Mapping[str, Any] | str | Path,
    verified: Mapping[str, Any] | str | Path,
    *,
    output_dir: str | Path | None = None,
    config_path: str | Path | None = None,
    model_names: Sequence[str] | None = None,
    targets: Sequence[str] | None = None,
    subject: str | None = None,
    section: str | None = None,
    subjects: Sequence[str] | None = None,
    benchmark_all: bool = False,
    question_numbers: Sequence[int] | None = None,
) -> dict[str, Path]:
    """@description 완료된 단일 실행 범위를 기존 공개 결과에 병합"""
    manifest = _as_exam(exam)
    run_data = _as_run(run)
    if run_data is None:
        raise ValueError("공개에는 실행 정본이 필요합니다.")
    if run_data.get("exam_id") != manifest.id:
        raise ValueError("시험과 실행 identity가 일치하지 않습니다.")
    mode_id = run_data.get("mode")
    mode = next((item for item in manifest.modes if item.id == mode_id), None)
    if mode is None or run_data.get("input_mode") != mode.input_mode:
        raise ValueError("실행 mode와 input_mode가 시험 매니페스트와 다릅니다.")
    sections, models = _selected_scope(
        manifest,
        run_data,
        model_names=model_names,
        targets=targets,
        subject=subject,
        section=section,
        subjects=subjects,
        benchmark_all=benchmark_all,
        question_numbers=question_numbers,
    )
    expected_numbers: dict[str, set[int]] = {}
    for item in sections:
        actual = set(_stored_question_numbers(run_data, item.target) or [])
        full = {question.number for question in load_section_questions(manifest, item)}
        if actual != full:
            raise ValueError(f"{item.target}는 전체 문항이 선택되지 않아 공개할 수 없습니다.")
        expected_numbers[item.target] = full
    verified_data = load_verified(verified) if isinstance(verified, (str, Path)) else copy.deepcopy(dict(verified))
    if verified_data.get("exam_id") not in (None, manifest.id):
        raise ValueError("검증 결과와 시험 identity가 일치하지 않습니다.")
    if verified_data.get("mode") not in (None, mode.id):
        raise ValueError("검증 결과와 실행 mode가 일치하지 않습니다.")
    rows_by_scope: dict[tuple[str, str], list[Mapping[str, Any]]] = {}
    for row in verified_data.get("results", []):
        if not isinstance(row, Mapping):
            continue
        target = row.get("target")
        model = row.get("model_name")
        if target in expected_numbers and model in models:
            rows_by_scope.setdefault((target, model), []).append(row)
    expected_scopes = {(item.target, model) for item in sections for model in models}
    if set(rows_by_scope) != expected_scopes:
        missing = sorted(expected_scopes - set(rows_by_scope))
        raise ValueError(f"선택한 섹션·모델의 검증 결과가 누락되었습니다: {missing}")
    for (target, model), scope_rows in rows_by_scope.items():
        numbers = {int(row["question_number"]) for row in scope_rows}
        if numbers != expected_numbers[target] or len(numbers) != len(scope_rows):
            raise ValueError(f"검증 결과의 문항 범위가 실행 입력과 다릅니다: {target}/{model}")
        if any(row.get("complete") is not True or row.get("needs_manual_review") is True for row in scope_rows):
            raise ValueError(f"미완료 또는 수동 검토 대상 결과는 공개할 수 없습니다: {target}/{model}")
    raw_by_key = {result_identity(result): result for result in run_data.get("results", [])}
    for (target, model), scope_rows in rows_by_scope.items():
        for row in scope_rows:
            if row.get("provenance") != "graded":
                continue
            raw_number = 0 if run_data["input_mode"] == "section" else int(row["question_number"])
            source = raw_by_key.get((target, model, raw_number))
            if (
                source is None
                or not result_is_completed(source)
                or row.get("raw_response", "") != source.get("raw_response", "")
                or (
                    "timestamp" in row
                    and row.get("timestamp") != source.get("timestamp")
                )
            ):
                raise ValueError(f"정본 원본이 없는 stale graded 결과는 공개할 수 없습니다: {target}/{model}")
    mode_dir = (
        Path(output_dir).expanduser().resolve()
        if output_dir is not None
        else manifest.project_root / "published" / manifest.id / mode.id
    )
    results_path = mode_dir / "results.json"
    token_path = mode_dir / "token_usage.json"
    metadata_path = mode_dir / "questions_metadata.json"
    public_records = _load_json(results_path, [])
    if not isinstance(public_records, list):
        raise ValueError("기존 published/results.json은 목록이어야 합니다.")
    existing: dict[tuple[str, str], Mapping[str, Any]] = {}
    for record in public_records:
        if not isinstance(record, Mapping) or not record.get("model_name"):
            continue
        canonical_target = _record_target(record, manifest)
        if canonical_target is not None:
            existing[(canonical_target, record["model_name"])] = record
    snapshots = {model["name"]: model for model in run_data["selected_models"]}
    changed_scopes: set[tuple[str, str]] = set()
    for identity, scope_rows in rows_by_scope.items():
        imported_only = all(row.get("provenance") == "imported_public" for row in scope_rows)
        if imported_only and identity in existing:
            continue
        target, model = identity
        existing[identity] = _public_record(scope_rows, manifest.section(target), snapshots.get(model))
        changed_scopes.add(identity)
    token_data = _load_json(token_path, {"models": {}})
    if not isinstance(token_data, Mapping):
        raise ValueError("기존 token_usage.json 형식이 잘못되었습니다.")
    token_updates: dict[tuple[str, str], Mapping[str, Any]] = {}
    for target, model in changed_scopes:
        section_key = _token_section_key(target, run_data["input_mode"])
        scope_rows = rows_by_scope[(target, model)]
        legacy_scope = all(
            row.get("provenance") in {"imported_public", "excel_manual"}
            for row in scope_rows
        )
        token_models = token_data.get("models", {})
        existing_model = token_models.get(model, {}) if isinstance(token_models, Mapping) else {}
        existing_sections = (
            existing_model.get("sections", {})
            if isinstance(existing_model, Mapping)
            else {}
        )
        if legacy_scope and isinstance(existing_sections, Mapping) and section_key in existing_sections:
            continue
        source_rows = [
            result
            for result in run_data.get("results", [])
            if result.get("target") == target
            and result.get("model_name") == model
            and result_is_completed(result)
        ]
        token_updates[(model, section_key)] = _new_token_record(source_rows)
    merged_tokens = _merge_token_scope(token_data, token_updates)
    metadata = _load_json(metadata_path, {})
    if not isinstance(metadata, Mapping):
        raise ValueError("기존 questions_metadata.json 형식이 잘못되었습니다.")
    merged_metadata = copy.deepcopy(dict(metadata))
    for item in sections:
        merged_metadata[_metadata_section_key(item)] = {
            str(question.number): {
                "hasImage": bool(question.image_paths),
                "points": int(question.points),
            }
            for question in load_section_questions(manifest, item)
        }
    metadata_models: Sequence[Mapping[str, Any]] = run_data["selected_models"]
    if config_path is not None:
        current_config = load_config(config_path, model_names=models, resolve_secrets=False)
        metadata_models = current_config["models"]
    model_metadata_path = manifest.project_root / "web" / "model_metadata.json"
    existing_model_metadata = _load_json(model_metadata_path, {})
    if not changed_scopes and config_path is None:
        merged_model_metadata = existing_model_metadata
    else:
        merged_model_metadata = merge_model_metadata(existing_model_metadata, metadata_models)
    _write_json_if_changed(results_path, list(existing.values()))
    _write_json_if_changed(token_path, merged_tokens)
    _write_json_if_changed(metadata_path, merged_metadata)
    _write_json_if_changed(model_metadata_path, merged_model_metadata)
    return {
        "results": results_path,
        "token_usage": token_path,
        "questions_metadata": metadata_path,
        "model_metadata": model_metadata_path,
    }


__all__ = [
    "export_run_to_excel",
    "import_excel_corrections",
    "publish_run",
    "target_to_section_key",
]
