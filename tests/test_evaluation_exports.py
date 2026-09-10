"""단일 실행 검증 결과의 공개·Excel 변환 집중 검증."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from csat_benchmark.evaluation import grade_run, save_verified
from csat_benchmark.exams import load_exam
from csat_benchmark.exports import (
    _merge_token_scope,
    _new_token_record,
    export_run_to_excel,
    import_excel_corrections,
    publish_run,
)
from csat_benchmark.runs import canonical_results_path, create_run, save_run


class _SectionVerifier:
    """섹션 응답에 고정된 문항별 답을 반환하는 검증기 대역."""

    model_id = "test-verifier"
    reasoning_effort = "low"

    def verify_hard_answers(self, raw_response, question_infos):
        answers = [int(value) for value in raw_response.split(",")]
        return [
            {
                "question_number": info["number"],
                "correct_answer": info["correct_answer"],
                "llm_answer": answers[index],
            }
            for index, info in enumerate(question_infos)
        ]


class _QuestionVerifier:
    """문항별 응답을 정수 답으로 변환하는 검증기 대역."""

    model_id = "test-verifier"
    reasoning_effort = "low"

    def verify_answer(self, raw_response, correct_answer, question_number, question_text):
        return int(raw_response)


def _make_exam(tmp_path: Path, *, mode: str = "default", input_mode: str = "section"):
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    questions = [
        {"number": 1, "correct_answer": 2, "points": 2, "question_path": "1.txt"},
        {"number": 2, "correct_answer": 1, "points": 2, "question_path": "2.txt"},
    ]
    (data_dir / "questions.json").write_text(
        json.dumps({"subject": "국어", "section": "예시", "questions": questions}, ensure_ascii=False),
        encoding="utf-8",
    )
    (data_dir / "1.txt").write_text("첫 번째 문항", encoding="utf-8")
    (data_dir / "2.txt").write_text("두 번째 문항", encoding="utf-8")
    manifest_path = tmp_path / "exam.json"
    manifest_path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "id": "exports-test",
                "title": "변환 시험",
                "short_name": "변환 시험",
                "exam_month": None,
                "data_dir": "data",
                "results_dir": "results",
                "publish": True,
                "sections": [
                    {
                        "target": "국어/예시",
                        "subject": "국어",
                        "section": "예시",
                        "group": "국어",
                        "kind": "common",
                        "questions": "questions.json",
                        "max_points": 4,
                    }
                ],
                "modes": [{"id": mode, "label": mode, "input_mode": input_mode}],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return load_exam(manifest_path, project_root=tmp_path)


def _make_run(exam, *, mode: str, input_mode: str, model: str = "모의 모델", raw: str = "2,1", success: bool = True):
    run = create_run(
        exam_id=exam.id,
        mode=mode,
        input_mode=input_mode,
        selected_targets=["국어/예시"],
        selected_models=[{"name": model, "price": {"input": 1.0}}],
        input_context=[{"target": "국어/예시", "question_numbers": [1, 2]}],
    )
    run["results"].append(
        {
            "target": "국어/예시",
            "subject": "국어",
            "section": "예시",
            "model_name": model,
            "question_number": 0 if input_mode == "section" else 1,
            "raw_response": raw if success else "",
            "timestamp": "2026-09-08T00:00:00+00:00",
            "success": success,
            "error_message": None if success else "기술 실패",
            "input_tokens": 5,
            "output_tokens": 7,
            "total_tokens": 12,
            "answer_status": None,
            "provider_stop_reason": None,
            "error_details": None,
        }
    )
    return run


def test_single_run_flat_public_and_excel_round_trip(tmp_path: Path):
    """단일 응답이 flat verified 행·모델 한 열·공개 결과로 변환된다."""
    exam = _make_exam(tmp_path)
    run = _make_run(exam, mode="default", input_mode="section")
    index = canonical_results_path(exam, "default")
    save_run(run, index)
    verified = grade_run(index, exam, _SectionVerifier())
    assert all("attempts" not in row for row in verified["results"])
    save_verified(verified, index)

    excel_path = export_run_to_excel(index, tmp_path / "answers.xlsx", run=index, exam=exam)
    worksheet = load_workbook(excel_path)["국어-예시"]
    assert [cell.value for cell in worksheet[1]] == ["문항 번호", "정답", "배점", "모의 모델"]
    worksheet["D2"] = 1
    worksheet.parent.save(excel_path)
    imported = import_excel_corrections(index, excel_path, run=index, exam=exam)
    assert imported["changed"] == 1
    assert imported["verified"]["results"][0]["provenance"] == "excel_manual"
    assert imported["verified"]["results"][0]["is_correct"] is False

    paths = publish_run(exam, index, index, output_dir=tmp_path / "published")
    public = json.loads(paths["results"].read_text(encoding="utf-8"))
    assert public[0]["score"] == 2
    assert "run_id" not in public[0]
    assert all("attempts" not in item for item in public[0]["results"])
    token = json.loads(paths["token_usage"].read_text(encoding="utf-8"))
    assert token["models"]["모의 모델"]["sections"]["국어-예시"]["total_tokens"] == 12
    assert token["models"]["모의 모델"]["sections"]["국어-예시"]["last_updated"] == "2026-09-08T00:00:00+00:00"
    assert token["models"]["모의 모델"]["last_updated"] == "2026-09-08T00:00:00+00:00"


def test_existing_workbook_preserves_unselected_cells_and_styles(tmp_path: Path):
    """기존 동적 헤더 Excel에서는 선택 열만 갱신한다."""
    exam = _make_exam(tmp_path)
    run = _make_run(exam, mode="default", input_mode="section")
    index = canonical_results_path(exam, "default")
    save_run(run, index)
    save_verified(grade_run(index, exam, _SectionVerifier()), index)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "국어-예시"
    worksheet["A1"] = "기존 제목"
    worksheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    worksheet["A2"] = "문항 번호"
    worksheet["B2"] = "정답"
    worksheet["C2"] = "기존 모델"
    worksheet["D2"] = "모의 모델"
    worksheet["A3"] = 1
    worksheet["B3"] = 2
    worksheet["C3"] = 9
    worksheet["D3"] = 9
    worksheet["A4"] = 2
    worksheet["B4"] = 1
    worksheet["C4"] = 8
    worksheet["D4"] = 8
    worksheet["A5"] = "총점"
    worksheet["C5"] = 17
    worksheet["D5"] = 17
    original_fill = worksheet["A1"].fill.fgColor.rgb
    excel_path = tmp_path / "existing.xlsx"
    workbook.save(excel_path)
    export_run_to_excel(index, excel_path, run=index, exam=exam, model_names=["모의 모델"])
    updated = load_workbook(excel_path)["국어-예시"]
    assert updated["A1"].value == "기존 제목"
    assert updated["A1"].fill.fgColor.rgb == original_fill
    assert updated["C3"].value == 9
    assert updated["D3"].value == 2


def test_excel_export_writes_totals_and_status_labels_without_replacing_formulas(tmp_path: Path):
    """새 Excel과 기존 literal 총점을 갱신하고 총점 수식은 보존한다."""
    exam = _make_exam(tmp_path)
    run = _make_run(exam, mode="default", input_mode="question")
    run["selected_models"].append({"name": "다른 모델", "price": {"input": 1.0}})
    verified = {
        "selected_models": ["모의 모델", "다른 모델"],
        "selected_targets": ["국어/예시"],
        "results": [
            {
                "target": "국어/예시",
                "model_name": "모의 모델",
                "question_number": 1,
                "extracted_answer": 2,
                "is_correct": True,
                "points": 2,
            },
            {
                "target": "국어/예시",
                "model_name": "모의 모델",
                "question_number": 2,
                "extracted_answer": -1,
                "is_correct": False,
                "points": 2,
            },
            {
                "target": "국어/예시",
                "model_name": "다른 모델",
                "question_number": 1,
                "extracted_answer": 1,
                "is_correct": False,
                "points": 2,
            },
            {
                "target": "국어/예시",
                "model_name": "다른 모델",
                "question_number": 2,
                "extracted_answer": -2,
                "is_correct": False,
                "points": 2,
            },
        ],
    }

    new_path = export_run_to_excel(verified, tmp_path / "new.xlsx", run=run, exam=exam)
    new_sheet = load_workbook(new_path, data_only=False)["국어-예시"]
    assert new_sheet["D2"].value == 2
    assert new_sheet["D3"].value == "(포기)"
    assert new_sheet["D3"].number_format == "General"
    assert new_sheet["D4"].value == 2
    assert new_sheet["E3"].value == "(검열)"
    assert new_sheet["E3"].number_format == "General"
    assert new_sheet["E4"].value == 0

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "국어-예시"
    worksheet.append(["문항 번호", "정답", "모의 모델", "다른 모델"])
    worksheet.append([1, 2, 9, 9])
    worksheet.append([2, 1, 8, 8])
    worksheet.append(["총점", None, "=SUM(C2:C3)", 99])
    existing_path = tmp_path / "existing_totals.xlsx"
    workbook.save(existing_path)

    export_run_to_excel(verified, existing_path, run=run, exam=exam)
    existing_sheet = load_workbook(existing_path, data_only=False)["국어-예시"]
    assert existing_sheet["C4"].value == "=SUM(C2:C3)"
    assert existing_sheet["D4"].value == 0


def test_imported_public_noop_preserves_public_and_tokens(tmp_path: Path):
    """raw 없는 imported_public 재공개는 기존 결과·토큰을 그대로 보존한다."""
    exam = _make_exam(tmp_path)
    run = _make_run(exam, mode="default", input_mode="section", success=False)
    index = canonical_results_path(exam, "default")
    save_run(run, index)
    imported_rows = [
        {
            "target": "국어/예시",
            "subject": "국어",
            "section": "예시",
            "model_name": "모의 모델",
            "question_number": 1,
            "extracted_answer": 2,
            "correct_answer": 2,
            "is_correct": True,
            "points": 2,
            "answer_status": "answered",
            "complete": True,
            "provenance": "imported_public",
        },
        {
            "target": "국어/예시",
            "subject": "국어",
            "section": "예시",
            "model_name": "모의 모델",
            "question_number": 2,
            "extracted_answer": 1,
            "correct_answer": 1,
            "is_correct": True,
            "points": 2,
            "answer_status": "answered",
            "complete": True,
            "provenance": "imported_public",
        },
    ]
    save_verified(
        {"schema_version": 1, "exam_id": exam.id, "mode": "default", "selected_models": ["모의 모델"], "results": imported_rows},
        index,
    )
    output = tmp_path / "published"
    output.mkdir()
    old_results = [{"sheet_name": "국어-예시", "subject": "국어", "section": "예시", "model_name": "모의 모델", "score": 4, "results": []}]
    old_tokens = {"models": {"모의 모델": {"total_tokens": 99, "sections": {"국어-예시": {"total_tokens": 99}}}}}
    results_path = output / "results.json"
    tokens_path = output / "token_usage.json"
    results_path.write_bytes((json.dumps(old_results, ensure_ascii=False, indent=2) + "\n").replace("\n", "\r\n").encode("utf-8"))
    tokens_path.write_bytes((json.dumps(old_tokens, ensure_ascii=False, indent=2) + "\n").replace("\n", "\r\n").encode("utf-8"))
    old_results_bytes = results_path.read_bytes()
    old_tokens_bytes = tokens_path.read_bytes()
    publish_run(exam, index, index, output_dir=output)
    assert results_path.read_bytes() == old_results_bytes
    assert tokens_path.read_bytes() == old_tokens_bytes


def test_incomplete_verified_scope_is_not_published(tmp_path: Path):
    """기술 실패로 남은 verified 행은 공개를 거부한다."""
    exam = _make_exam(tmp_path)
    run = _make_run(exam, mode="default", input_mode="section", success=False)
    index = canonical_results_path(exam, "default")
    save_run(run, index)
    save_verified(grade_run(index, exam, _SectionVerifier()), index)
    with pytest.raises(ValueError, match="미완료"):
        publish_run(exam, index, index, output_dir=tmp_path / "published")


def test_publish_preserves_legacy_tokens_and_marks_fresh_unknown_usage(tmp_path: Path):
    """Excel 수동 수정 범위의 기존 토큰을 보존하고 새 미상 사용량은 None으로 기록한다."""
    exam = _make_exam(tmp_path)
    legacy_model = "Claude Opus 4.7 (high)"
    fresh_model = "새 모델"
    run = _make_run(exam, mode="default", input_mode="section", model=legacy_model)
    run["selected_models"].append({"name": fresh_model, "price": {"input": 1.0}})
    run["results"][0].pop("input_tokens")
    run["results"][0].pop("output_tokens")
    run["results"][0].pop("total_tokens")
    run["results"].append(
        {
            "target": "국어/예시",
            "subject": "국어",
            "section": "예시",
            "model_name": fresh_model,
            "question_number": 0,
            "raw_response": "2,1",
            "timestamp": "2026-09-08T00:00:00+00:00",
            "success": True,
            "error_message": None,
            "answer_status": None,
            "provider_stop_reason": None,
            "error_details": None,
        }
    )
    index = canonical_results_path(exam, "default")
    save_run(run, index)
    verified_rows = []
    for model, provenance in ((legacy_model, "imported_public"), (fresh_model, "graded")):
        verified_rows.extend(
            {
                "target": "국어/예시",
                "subject": "국어",
                "section": "예시",
                "model_name": model,
                "question_number": number,
                "extracted_answer": answer,
                "correct_answer": correct,
                "is_correct": answer == correct,
                "points": 2,
                "answer_status": "answered",
                "provider_stop_reason": None,
                "needs_manual_review": False,
                "complete": True,
                "raw_response": "2,1" if provenance == "graded" else "",
                "timestamp": "2026-09-08T00:00:00+00:00" if provenance == "graded" else None,
                "provenance": provenance,
            }
            for number, answer, correct in ((1, 2, 2), (2, 1, 1))
        )
    save_verified(
        {
            "schema_version": 1,
            "exam_id": exam.id,
            "mode": "default",
            "selected_models": [legacy_model, fresh_model],
            "results": verified_rows,
        },
        index,
    )
    excel_path = export_run_to_excel(index, tmp_path / "answers.xlsx", run=index, exam=exam)
    workbook = load_workbook(excel_path)
    workbook["국어-예시"]["D2"] = 1
    workbook.save(excel_path)
    workbook.close()
    import_excel_corrections(index, excel_path, run=index, exam=exam)

    legacy_section = {
        "input_tokens": 101,
        "output_tokens": 202,
        "total_tokens": 303,
        "question_count": 2,
        "last_updated": "legacy-section",
        "custom": "keep",
    }
    unrelated_model = {
        "total_input_tokens": 7,
        "total_output_tokens": 8,
        "total_tokens": 15,
        "question_count": 1,
        "sections": {"다른 섹션": {"input_tokens": 7, "output_tokens": 8, "total_tokens": 15}},
    }
    old_tokens = {
        "models": {
            legacy_model: {
                "total_input_tokens": 101,
                "total_output_tokens": 202,
                "total_tokens": 303,
                "question_count": 2,
                "last_updated": "legacy-model",
                "sections": {"국어-예시": legacy_section},
            },
            "관계없는 모델": unrelated_model,
        }
    }
    output = tmp_path / "published"
    output.mkdir()
    (output / "results.json").write_text("[]\n", encoding="utf-8")
    (output / "token_usage.json").write_text(
        json.dumps(old_tokens, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    publish_run(exam, index, index, output_dir=output)
    tokens = json.loads((output / "token_usage.json").read_text(encoding="utf-8"))
    assert tokens["models"][legacy_model] == old_tokens["models"][legacy_model]
    assert tokens["models"]["관계없는 모델"] == unrelated_model
    assert tokens["models"][fresh_model]["sections"]["국어-예시"] == {
        "input_tokens": None,
        "output_tokens": None,
        "total_tokens": None,
        "question_count": 1,
        "last_updated": "2026-09-08T00:00:00+00:00",
    }
    assert tokens["models"][fresh_model]["total_input_tokens"] is None
    assert tokens["models"][fresh_model]["total_output_tokens"] is None
    assert tokens["models"][fresh_model]["total_tokens"] is None


def test_token_usage_dates_propagate_latest_section_and_preserve_legacy_date():
    """@description 여러 섹션의 최신 날짜를 모델에 반영하고 날짜 없는 원본의 기존 날짜 보존"""
    latest = "2026-09-08 20:52:10"
    older = "2026-09-07 20:52:10"
    source_record = _new_token_record(
        [
            {"input_tokens": 1, "output_tokens": 2, "total_tokens": 3, "timestamp": latest},
            {"input_tokens": 4, "output_tokens": 5, "total_tokens": 6, "timestamp": older},
        ]
    )
    legacy_record = {
        "input_tokens": 10,
        "output_tokens": 20,
        "total_tokens": 30,
        "question_count": 1,
        "last_updated": older,
    }
    merged = _merge_token_scope(
        {
            "models": {
                "모델": {
                    "last_updated": latest,
                    "sections": {
                        "최신 섹션": {
                            "input_tokens": 7,
                            "output_tokens": 8,
                            "total_tokens": 15,
                            "question_count": 1,
                            "last_updated": latest,
                        },
                        "기존 섹션": legacy_record,
                    },
                }
            }
        },
        {
            ("모델", "기존 섹션"): _new_token_record(
                [{"input_tokens": 11, "output_tokens": 22, "total_tokens": 33}]
            ),
            ("모델", "새 섹션"): source_record,
        },
    )

    model = merged["models"]["모델"]
    assert source_record["last_updated"] == latest
    assert model["sections"]["기존 섹션"]["last_updated"] == older
    assert model["sections"]["새 섹션"]["last_updated"] == latest
    assert model["last_updated"] == latest
