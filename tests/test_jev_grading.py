"""저장된 Jev 답안의 규칙 기반 채점 및 검증 결과 연동 검사."""

from __future__ import annotations

import json
from pathlib import Path
import pytest
from openpyxl import load_workbook

import jev_verify_answers
import verify_answers
from csat_benchmark.evaluation import EvaluationError, grade_run, load_verified, save_verified
from csat_benchmark.exams import load_exam
from csat_benchmark.exports import export_run_to_excel, publish_run
from csat_benchmark.jev_grading import verify_jev_section_result, verify_jev_single_result
from csat_benchmark.runs import canonical_results_path, create_run, save_run


def _make_exam(tmp_path: Path, *, input_mode: str):
    """세 자리 답안·복수 정답·선택형 문항이 있는 시험을 생성"""
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    (data_dir / "questions.json").write_text(
        json.dumps(
            {
                "subject": "수학",
                "section": "예시",
                "questions": [
                    {"number": 7, "correct_answer": 7, "points": 2, "question_text": "7번"},
                    {"number": 42, "correct_answer": [41, 42], "points": 3, "question_text": "42번"},
                    {"number": 43, "correct_answer": 4, "points": 5, "question_text": "43번"},
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    manifest_path = tmp_path / "exam.json"
    manifest_path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "id": f"jev-{input_mode}",
                "title": "Jev 채점 시험",
                "data_dir": "data",
                "results_dir": "results",
                "sections": [
                    {
                        "target": "수학/예시",
                        "subject": "수학",
                        "section": "예시",
                        "group": "수학",
                        "kind": "common",
                        "questions": "questions.json",
                        "max_points": 10,
                    }
                ],
                "modes": [{"id": "default", "label": "기본", "input_mode": input_mode}],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return load_exam(manifest_path, project_root=tmp_path)


def _jev_response() -> str:
    """Jev의 자리별 정수 응답과 선택 응답을 JSON 문자열로 반환"""
    return json.dumps(
        {
            "answers": {
                "q7_hundreds": {"type": "choice", "choice": "0"},
                "q7_tens": {"type": "choice", "choice": "0"},
                "q7_units": {"type": "choice", "choice": "7"},
                "q42_hundreds": {"type": "choice", "choice": "0"},
                "q42_tens": {"type": "choice", "choice": "4"},
                "q42_units": {"type": "choice", "choice": "2"},
                "q43": {"type": "choice", "choice": "3"},
            }
        },
        ensure_ascii=False,
    )


def _make_run(exam, *, input_mode: str, ordinary_model: bool = False):
    """Jev 모델과 선택적 기존 모델 결과를 가진 저장 실행 생성"""
    models = [{"name": "jev-model", "api_type": "jev", "model_id": "jev"}]
    if ordinary_model:
        models.append({"name": "ordinary-model", "api_type": "openai", "model_id": "ordinary"})
    run = create_run(
        exam_id=exam.id,
        mode="default",
        input_mode=input_mode,
        selected_targets=["수학/예시"],
        selected_models=models,
        input_context=[{"target": "수학/예시", "question_numbers": [7, 42, 43]}],
    )
    run["results"] = []
    question_numbers = [0] if input_mode == "section" else [7, 42, 43]
    for model in models:
        for number in question_numbers:
            is_jev = model["api_type"] == "jev"
            run["results"].append(
                {
                    "target": "수학/예시",
                    "subject": "수학",
                    "section": "예시",
                    "model_name": model["name"],
                    "question_number": number,
                    "raw_response": _jev_response() if is_jev else "ordinary raw response",
                    "timestamp": f"2026-09-28T00:00:{number:02d}+00:00",
                    "success": True,
                    "answer_status": None,
                    "provider_stop_reason": None,
                    "input_tokens": 11,
                    "output_tokens": 13,
                    "total_tokens": 24,
                }
            )
    return run


def _prior_row(
    model_name: str,
    number: int,
    answer: int,
    correct_answer: int | list[int],
    *,
    raw_response: str = "previous response",
) -> dict:
    """기존 완료 verified 행 생성"""
    return {
        "target": "수학/예시",
        "subject": "수학",
        "section": "예시",
        "model_name": model_name,
        "question_number": number,
        "extracted_answer": answer,
        "correct_answer": correct_answer,
        "is_correct": answer in correct_answer if isinstance(correct_answer, list) else answer == correct_answer,
        "points": {7: 2, 42: 3, 43: 5}[number],
        "answer_status": "answered",
        "complete": True,
        "raw_response": raw_response,
        "provenance": "graded",
    }


def test_section_grading_decodes_once_and_respects_subset(tmp_path: Path, monkeypatch):
    """섹션 응답 한 번 파싱, 앞자리 0, 복수 정답과 선택 문항 범위를 확인"""
    import csat_benchmark.jev_grading as jev_grading

    exam = _make_exam(tmp_path, input_mode="section")
    run = _make_run(exam, input_mode="section", ordinary_model=True)
    raw_response = run["results"][0]["raw_response"]
    load_calls = 0
    decoded_numbers = []
    original_loader = jev_grading._load_jev_answers
    original_decoder = jev_grading.decode_jev_answer

    def _counted_loader(result):
        nonlocal load_calls
        load_calls += 1
        return original_loader(result)

    def _counted_decoder(answers, number):
        decoded_numbers.append(number)
        return original_decoder(answers, number)

    monkeypatch.setattr(jev_grading, "_load_jev_answers", _counted_loader)
    monkeypatch.setattr(jev_grading, "decode_jev_answer", _counted_decoder)
    full = grade_run(
        run,
        exam,
        None,
        default_model_names=["jev-model"],
        section_result_extractor=verify_jev_section_result,
        grading_description="Jev 저장 응답 직접 채점",
    )
    assert load_calls == 1
    assert full["selected_models"] == ["jev-model"]
    assert {row["question_number"]: row["extracted_answer"] for row in full["results"]} == {
        7: 7,
        42: 42,
        43: 3,
    }
    assert {row["question_number"]: row["is_correct"] for row in full["results"]} == {
        7: True,
        42: True,
        43: False,
    }
    assert full["score_by_model"]["jev-model"] == 5
    assert all(row["answer_status"] == "answered" for row in full["results"])
    assert all(row["needs_manual_review"] is False for row in full["results"])
    assert all(row["raw_response"] == raw_response for row in full["results"])
    assert all(row["timestamp"].startswith("2026-09-28T00:00:") for row in full["results"])
    assert decoded_numbers == [7, 42, 43]

    load_calls = 0
    decoded_numbers.clear()
    subset = grade_run(
        run,
        exam,
        None,
        default_model_names=["jev-model"],
        section_result_extractor=verify_jev_section_result,
        verified={
            "selected_models": ["jev-model"],
            "results": [
                _prior_row("jev-model", 7, 7, 7),
                _prior_row("jev-model", 43, 3, 4),
            ],
        },
        grading_description="Jev 저장 응답 직접 채점",
    )
    assert load_calls == 1
    assert decoded_numbers == [42]
    subset_by_number = {row["question_number"]: row for row in subset["results"]}
    assert subset_by_number[42]["extracted_answer"] == 42
    assert subset["score_by_model"]["jev-model"] == 5


def test_question_mode_grades_each_raw_answer_without_verifier(tmp_path: Path, monkeypatch):
    """문항별 Jev 응답도 API 키 없이 한 번씩 결정적으로 채점"""
    import csat_benchmark.jev_grading as jev_grading

    exam = _make_exam(tmp_path, input_mode="question")
    run = _make_run(exam, input_mode="question")
    calls = 0
    original_loader = jev_grading._load_jev_answers

    def _counted_loader(result):
        nonlocal calls
        calls += 1
        return original_loader(result)

    monkeypatch.setattr(jev_grading, "_load_jev_answers", _counted_loader)
    result = grade_run(
        run,
        exam,
        None,
        single_result_extractor=verify_jev_single_result,
        grading_description="Jev 저장 응답 직접 채점",
    )
    assert calls == 3
    assert {row["question_number"] for row in result["results"]} == {7, 42, 43}
    assert result["score_by_model"]["jev-model"] == 5


def test_missing_saved_fields_raise_and_technical_failure_stays_incomplete(tmp_path: Path, monkeypatch):
    """손상된 성공 응답은 오류로 알리고 전송 실패는 미완료 행으로 남김"""
    import csat_benchmark.jev_grading as jev_grading

    with pytest.raises(EvaluationError, match="answers 객체"):
        verify_jev_single_result(
            None,
            {"model_name": "jev-model", "raw_response": "{}"},
            {"number": 7, "correct_answer": 7, "points": 2},
            1,
            Path.cwd(),
        )
    with pytest.raises(EvaluationError, match="자리별 응답"):
        verify_jev_single_result(
            None,
            {
                "model_name": "jev-model",
                "raw_response": json.dumps(
                    {"answers": {"q7_units": {"type": "choice", "choice": "7"}}}
                ),
            },
            {"number": 7, "correct_answer": 7, "points": 2},
            1,
            Path.cwd(),
        )

    exam = _make_exam(tmp_path, input_mode="section")
    run = _make_run(exam, input_mode="section")
    source = run["results"][0]
    source.update(
        success=False,
        raw_response="",
        answer_status="technical_failure",
        error_message="network error",
    )
    monkeypatch.setattr(
        jev_grading,
        "decode_jev_answer",
        lambda *_args: pytest.fail("기술 실패 응답을 디코드하면 안 됩니다."),
    )
    result = grade_run(
        run,
        exam,
        None,
        default_model_names=["jev-model"],
        section_result_extractor=verify_jev_section_result,
        grading_description="Jev 저장 응답 직접 채점",
    )
    assert all(row["complete"] is False for row in result["results"])
    assert all(row["is_correct"] is None for row in result["results"])
    assert all(row["answer_status"] == "incomplete" for row in result["results"])
    assert all(row["error_message"] == "network error" for row in result["results"])


def test_jev_cli_ignores_config_preserves_other_sidecar_regrades_and_exports(
    tmp_path: Path, monkeypatch
):
    """Jev 전용 선택·재채점·기존 모델 보존·Excel 내보내기를 확인"""
    exam = _make_exam(tmp_path, input_mode="section")
    run = _make_run(exam, input_mode="section", ordinary_model=True)
    index = canonical_results_path(exam, "default")
    save_run(run, index)
    wrong_jev_rows = [
        _prior_row("jev-model", 7, 99, 7),
        _prior_row("jev-model", 42, 99, [41, 42]),
        _prior_row("jev-model", 43, 99, 4),
    ]
    ordinary_rows = [
        _prior_row("ordinary-model", 7, 90, 7, raw_response="ordinary raw response"),
        _prior_row("ordinary-model", 42, 91, [41, 42], raw_response="ordinary raw response"),
        _prior_row("ordinary-model", 43, 92, 4, raw_response="ordinary raw response"),
    ]
    save_verified(
        {
            "schema_version": 1,
            "exam_id": exam.id,
            "mode": "default",
            "selected_models": ["jev-model", "ordinary-model"],
            "results": wrong_jev_rows + ordinary_rows,
        },
        index,
    )
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.setattr(
        verify_answers,
        "build_verifier",
        lambda *_args: pytest.fail("Jev grading must not build an LLM verifier."),
    )

    assert jev_verify_answers.main(["--exam", str(exam.manifest_path), "--config", "missing.json"]) == 0
    unchanged = load_verified(index)
    by_model_number = {
        (row["model_name"], row["question_number"]): row for row in unchanged["results"]
    }
    assert [by_model_number[("jev-model", number)]["extracted_answer"] for number in (7, 42, 43)] == [99] * 3
    assert [by_model_number[("ordinary-model", number)]["extracted_answer"] for number in (7, 42, 43)] == [90, 91, 92]

    assert jev_verify_answers.main(
        ["--exam", str(exam.manifest_path), "--config", "missing.json", "--update"]
    ) == 0
    updated = load_verified(index)
    by_model_number = {(row["model_name"], row["question_number"]): row for row in updated["results"]}
    assert [by_model_number[("jev-model", number)]["extracted_answer"] for number in (7, 42, 43)] == [7, 42, 3]
    assert [by_model_number[("ordinary-model", number)]["extracted_answer"] for number in (7, 42, 43)] == [90, 91, 92]
    assert all(
        by_model_number[("jev-model", number)]["raw_response"] == _jev_response()
        and by_model_number[("jev-model", number)]["timestamp"] == "2026-09-28T00:00:00+00:00"
        for number in (7, 42, 43)
    )

    save_verified(
        {
            "schema_version": 1,
            "exam_id": exam.id,
            "mode": "default",
            "selected_models": ["jev-model"],
            "results": wrong_jev_rows,
        },
        index,
    )
    assert jev_verify_answers.main(
        ["--exam", str(exam.manifest_path), "--config", "missing.json", "--models", "jev-model"]
    ) == 0
    explicit_regrade = load_verified(index)
    by_model_number = {
        (row["model_name"], row["question_number"]): row for row in explicit_regrade["results"]
    }
    assert [by_model_number[("jev-model", number)]["extracted_answer"] for number in (7, 42, 43)] == [7, 42, 3]
    assert [by_model_number[("ordinary-model", number)]["extracted_answer"] for number in (7, 42, 43)] == [90, 91, 92]

    excel_path = export_run_to_excel(index, tmp_path / "jev.xlsx", run=index, exam=exam)
    worksheet = load_workbook(excel_path)["수학-예시"]
    header_to_column = {
        worksheet.cell(1, column).value: column
        for column in range(1, worksheet.max_column + 1)
    }
    rows_by_number = {
        worksheet.cell(row, 1).value: row
        for row in range(2, worksheet.max_row + 1)
        if isinstance(worksheet.cell(row, 1).value, int)
    }
    assert worksheet.cell(rows_by_number[7], header_to_column["jev-model"]).value == 7
    assert worksheet.cell(rows_by_number[42], header_to_column["jev-model"]).value == 42

    published_paths = publish_run(
        exam,
        index,
        index,
        output_dir=tmp_path / "published",
    )
    published = json.loads(published_paths["results"].read_text(encoding="utf-8"))
    jev_record = next(record for record in published if record["model_name"] == "jev-model")
    assert jev_record["score"] == 5
    assert jev_record["total_questions"] == 3
    assert [row["extracted_answer"] for row in jev_record["results"]] == [7, 42, 3]
