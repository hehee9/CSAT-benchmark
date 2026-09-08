"""@brief 공개 예제를 로컬 모의 API로 실행·채점하고 Excel·공개 JSON까지 확인한다."""

from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import load_workbook

_PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

from csat_benchmark.runs import load_run  # noqa: E402
from csat_benchmark.evaluation import load_verified  # noqa: E402


class _MockAPI(BaseHTTPRequestHandler):
    """@brief 외부 연결 없이 공개 예제의 생성·답안 추출 응답을 제공한다."""

    generations: list[dict] = []
    verifications: list[dict] = []
    lock = threading.Lock()

    def log_message(self, format, *args):
        """@brief 모의 HTTP 접근 로그를 생략한다."""

    def do_POST(self):
        """@brief Chat 스트림과 Responses 답안 추출 요청에 응답한다."""
        body = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
        if self.path == '/v1/chat/completions':
            with self.lock:
                self.generations.append(body)
                request_number = len(self.generations)
            input_tokens = 100
            output_tokens = 20
            chunks = [
                {'choices': [{'index': 0, 'delta': {'content': '정답은 ②입니다.'}, 'finish_reason': None}]},
                {'choices': [], 'usage': {'prompt_tokens': input_tokens,
                                         'completion_tokens': output_tokens,
                                         'total_tokens': input_tokens + output_tokens}},
            ]
            payload = ''.join(
                'data: ' + json.dumps({'id': f'chatcmpl-mock-{request_number}',
                                     'object': 'chat.completion.chunk', 'created': 0,
                                     'model': body['model'], **chunk}, ensure_ascii=False) + '\n\n'
                for chunk in chunks
            ) + 'data: [DONE]\n\n'
            content_type = 'text/event-stream'
        elif self.path == '/v1/responses':
            self.verifications.append(body)
            properties = body['text']['format']['schema']['properties']
            answer = json.dumps(
                {'answers': [{'question_number': number, 'correct_answer': 2, 'llm_answer': 2}
                             for number in (1, 2)]}
                if 'answers' in properties else {'correct_answer': 2, 'llm_answer': 2}
            )
            payload = json.dumps({
                'id': 'resp-mock', 'object': 'response', 'created_at': 0,
                'status': 'completed', 'model': body['model'],
                'output': [{'id': 'msg-mock', 'type': 'message', 'role': 'assistant',
                            'status': 'completed',
                            'content': [{'type': 'output_text', 'text': answer, 'annotations': []}]}],
            })
            content_type = 'application/json'
        else:
            self.send_error(404)
            return
        encoded = payload.encode('utf-8')
        self.send_response(200)
        self.send_header('Content-Type', content_type)
        self.send_header('Content-Length', str(len(encoded)))
        self.end_headers()
        self.wfile.write(encoded)


def _run(root: Path, *args: str) -> str:
    """@brief 현재 Python 환경에서 실제 공개 CLI를 실행한다."""
    result = subprocess.run([sys.executable, *args], cwd=root, text=True,
                            encoding='utf-8', capture_output=True, check=True)
    print(result.stdout, end='')
    return result.stdout


def main():
    """@brief 임시 문항으로 일반·쉬움 단일 실행 명령 연결을 확인한다."""
    root = _PROJECT_ROOT
    os.environ['PYTHONIOENCODING'] = 'utf-8'
    os.environ['CSAT_MOCK_API_KEY'] = 'local-mock-only'
    server = ThreadingHTTPServer(('127.0.0.1', 0), _MockAPI)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        with tempfile.TemporaryDirectory(prefix='csat-public-workflow-') as temporary:
            temporary_root = Path(temporary)
            config_path = temporary_root / 'config.json'
            base_url = f'http://127.0.0.1:{server.server_port}/v1'
            config_path.write_text(json.dumps({
                'system_prompt': '공개 예제 모의 검증',
                'models': [{'name': '모의 모델', 'api_type': 'openai',
                            'api_key_env': 'CSAT_MOCK_API_KEY', 'model_id': 'mock',
                            'request_api': 'chat_completions', 'base_url': base_url,
                            'price': {'input': 1.0, 'output': 2.0}}],
                'verifier': {'api_key_env': 'CSAT_MOCK_API_KEY', 'model_id': 'mock-verifier',
                             'base_url': base_url},
            }, ensure_ascii=False), encoding='utf-8')
            (temporary_root / 'questions.json').write_text(json.dumps({
                'subject': '국어', 'section': '예시',
                'questions': [{'number': number, 'correct_answer': 2, 'points': 2,
                               'question_text': f'{number}번 모의 문항: 1+1은?'}
                              for number in (1, 2)],
            }, ensure_ascii=False), encoding='utf-8')
            manifest = temporary_root / 'exam.json'
            manifest.write_text(json.dumps({
                'schema_version': 1, 'id': 'workflow-test', 'title': '명령 연결 확인',
                'data_dir': '.', 'results_dir': 'results', 'publish': False,
                'sections': [{'target': '국어/예시', 'subject': '국어', 'section': '예시',
                              'group': '국어', 'kind': 'common',
                              'questions': 'questions.json', 'max_points': 4}],
                'modes': [{'id': 'default', 'label': '일반', 'input_mode': 'section'},
                          {'id': 'easy', 'label': '쉬움', 'input_mode': 'question'}],
            }, ensure_ascii=False), encoding='utf-8')
            common = ['--exam', str(manifest), '--config', str(config_path)]
            for mode, flags, expected_calls in [('default', [], 1), ('easy', ['--easy'], 2)]:
                generation_start = len(_MockAPI.generations)
                verification_start = len(_MockAPI.verifications)
                _run(root, 'api_solver.py', *common, *flags, '--check')
                _run(root, 'api_solver.py', *common, *flags)
                _run(root, 'verify_answers.py', *common, *flags)
                index = temporary_root / 'results' / mode / 'results.json'
                raw_document = json.loads(index.read_text(encoding='utf-8'))
                assert raw_document['schema_version'] == 3
                assert 'model_results' in raw_document and 'results' not in raw_document
                raw = load_run(index)
                verified = load_verified(index)
                assert len(_MockAPI.generations) - generation_start == expected_calls
                assert len(_MockAPI.verifications) - verification_start == expected_calls * 2
                assert len(raw['results']) == expected_calls
                assert all(row['success'] and row['raw_response'] for row in raw['results'])
                assert all('attempt' not in row for row in raw['results'])
                assert verified['score_by_model']['모의 모델'] == 4
                excel = index.parent / 'answers.xlsx'
                _run(root, 'sync_data.py', 'import', *common, *flags, '--excel', str(excel))
                workbook = load_workbook(excel)
                sheet = workbook['국어-예시']
                headers = [cell.value for cell in sheet[1]]
                assert headers == ['문항 번호', '정답', '배점', '모의 모델']
                assert sheet.cell(4, 4).value == 4
                sheet.cell(2, 4).value = 1
                workbook.save(excel)
                workbook.close()
                _run(root, 'sync_data.py', 'export', *common, *flags, '--excel', str(excel))
                _run(root, 'sync_data.py', 'publish', *common, *flags)
                public_dir = temporary_root / 'published/workflow-test' / mode
                public = json.loads((public_dir / 'results.json').read_text(encoding='utf-8'))
                usage = json.loads((public_dir / 'token_usage.json').read_text(encoding='utf-8'))
                model = usage['models']['모의 모델']
                assert public[0]['score'] == 2
                assert [row['is_correct'] for row in public[0]['results']] == [False, True]
                assert model['total_input_tokens'] == 100 * expected_calls
                assert model['total_output_tokens'] == 20 * expected_calls
                for filename in ('results.json', 'token_usage.json', 'questions_metadata.json'):
                    text = (public_dir / filename).read_text(encoding='utf-8')
                    assert all(secret not in text for secret in (
                        'raw_response', 'question_text', 'api_key', 'local-mock-only', 'attempt', 'run_id',
                    ))
                assert load_run(index) == raw
                if mode == 'easy':
                    _run(root, 'api_solver.py', *common, *flags, '--question-numbers', '1')
                    assert [row['question_number'] for row in load_verified(index)['results']] == [2]
                    _run(root, 'verify_answers.py', *common, *flags, '--question-numbers', '1')
                    assert len(load_verified(index)['results']) == 2
                    _run(root, 'sync_data.py', 'publish', *common, *flags)
                    updated = json.loads((public_dir / 'results.json').read_text(encoding='utf-8'))
                    assert len(updated) == 1 and updated[0]['score'] == 4
                print(f'단일 실행 통합 검증 통과: {mode}')
    finally:
        server.shutdown()
        server.server_close()
        thread.join()


if __name__ == '__main__':
    main()
